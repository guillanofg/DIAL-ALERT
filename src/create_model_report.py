"""Create the DIAL-ALERT Step 4 Model Implementation Report."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from create_eda_report import (
    ARTIFACTS,
    MID_GRAY,
    add_body,
    add_bullets,
    add_code,
    add_figure,
    add_page_break,
    add_table,
    configure_document,
    set_repeat_page_number,
    set_font,
)


ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
MODELS = ROOT / "models"
REPORTS.mkdir(exist_ok=True)
OUTPUT = REPORTS / "Franklin_Guillano_DIAL_ALERT_Model_Implementation_Report.docx"


def fmt_ci(metric_row: pd.Series, digits: int = 3) -> str:
    return (
        f"{metric_row['estimate']:.{digits}f} "
        f"({metric_row['ci_lower_95']:.{digits}f} to {metric_row['ci_upper_95']:.{digits}f})"
    )


def make_report() -> Path:
    comparison = pd.read_csv(ARTIFACTS / "model_comparison.csv")
    split = pd.read_csv(ARTIFACTS / "split_summary.csv")
    calibration = pd.read_csv(ARTIFACTS / "calibration_method_comparison.csv")
    intervals = pd.read_csv(ARTIFACTS / "test_metric_confidence_intervals.csv").set_index("metric")
    capacity = pd.read_csv(ARTIFACTS / "capacity_metrics.csv")
    importance = pd.read_csv(ARTIFACTS / "permutation_importance.csv")
    final = json.loads((ARTIFACTS / "final_test_metrics.json").read_text(encoding="utf-8"))
    manifest = json.loads((MODELS / "model_manifest.json").read_text(encoding="utf-8"))
    metrics = final["test_metrics_calibrated"]
    selected = final["selected_model"]
    threshold = float(final["operating_threshold"])

    doc = Document()
    configure_document(doc)
    footer = doc.sections[0].footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("DIAL ALERT   |   Model Implementation Report   |   "), size=8.5, color=MID_GRAY)
    set_repeat_page_number(footer)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_before = Pt(90)
    title.add_run("DIAL ALERT Model Implementation Report")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("Capstone Project Step 4")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run("Patient grouped model comparison for early prediction of blood pressure defined intradialytic hypotension")
    set_font(r, size=12, bold=True)
    p = doc.add_paragraph()
    set_font(p.add_run("Prepared by Franklin Guillano"), size=11)
    p = doc.add_paragraph()
    set_font(p.add_run("September 2026"), size=10.5, color=MID_GRAY)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(105)
    set_font(
        p.add_run(
            "Academic clinical decision support prototype. External and prospective validation are required before clinical use."
        ),
        size=9.5,
        italic=True,
        color=MID_GRAY,
    )

    add_page_break(doc)
    doc.add_heading("Executive Summary", level=1)
    add_body(
        doc,
        f"Seven modelling configurations were compared using patient-grouped cross-validation and a patient-disjoint validation set. {selected} was selected under the locked rule: maximize grouped cross-validation average precision, then among models within 0.01 of the best result choose the model with the lower validation Brier score. This rule balanced event ranking with probability reliability and selected the random forest over the less well-calibrated boosted-tree candidate."
    )
    add_body(
        doc,
        f"On the untouched test set of {int(split.loc[split.split.eq('test'), 'sessions'].iloc[0]):,} sessions, the selected model achieved average precision {metrics['average_precision']:.3f}, ROC AUC {metrics['roc_auc']:.3f}, and Brier score {metrics['brier_score']:.3f}. At the validation-selected F2 threshold of {threshold:.3f}, sensitivity was {metrics['sensitivity']:.1%}, specificity {metrics['specificity']:.1%}, precision {metrics['precision']:.1%}, and F1 score {metrics['f1']:.3f}. Patient-cluster bootstrap intervals quantify uncertainty from repeated sessions within patients."
    )
    add_body(
        doc,
        f"If operational capacity limits alerts to the highest-risk 20% of sessions, DIAL-ALERT captures {metrics['recall_at_capacity']:.1%} of observed events with {metrics['precision_at_capacity']:.1%} positive predictive value and {metrics['lift_at_capacity']:.2f}-fold lift over prevalence. These results support continued internal development, not clinical deployment."
    )

    doc.add_heading("Rubric Evidence", level=2)
    rubric_rows = [
        ["Multiple models", "Dummy, logistic, L1-selected logistic, decision tree, PCA logistic, random forest, and histogram gradient boosting were compared."],
        ["Appropriate tuning", "Grid or randomized searches were nested within three-fold patient-grouped cross-validation."],
        ["Relevant metrics", "Average precision is primary; ROC AUC, Brier score, log loss, sensitivity, specificity, precision, F1, and alert-capacity metrics are reported."],
        ["Fair comparison", "Every candidate uses the same cohort, feature boundary, patient splits, and outcome."],
        ["Reproducibility", "Configurations, split assignments, candidate pipelines, final predictor, threshold, package versions, hashes, and random seeds are saved."],
        ["Model choice", "The selection rule was locked before final testing and is explained using discrimination, calibration, and operational needs."],
    ]
    add_table(doc, ["Requirement", "Evidence"], rubric_rows, widths=[1.55, 5.35], font_size=8.6)

    add_page_break(doc)
    doc.add_heading("1 Prediction Task", level=1)
    add_body(
        doc,
        "DIAL-ALERT is a supervised binary-classification task. For each eligible haemodialysis session, the model estimates the probability of a later systolic blood-pressure measurement below 90 mmHg using 22 fields available at the index time or from earlier sessions. The target is operationally useful for risk ranking but is described as BP-defined IDH because symptoms and interventions are unavailable in HEMOBP."
    )
    task_rows = [
        ["Analysis unit", "One eligible patient-day haemodialysis session"],
        ["Prediction time", "Earliest valid active-dialysis record during minutes 0 to 30"],
        ["Target", "Any later SBP below 90 mmHg"],
        ["Primary metric", "Average precision because event prevalence is 8.49%"],
        ["Secondary metrics", "ROC AUC, Brier score, log loss, sensitivity, specificity, precision, and F1"],
        ["Operational measure", "Recall, precision, lift, and false-alert burden at fixed alert capacity"],
    ]
    add_table(doc, ["Element", "Definition"], task_rows, widths=[1.7, 5.2], font_size=9.1)
    add_figure(
        doc,
        ARTIFACTS / "step4_model_workflow.png",
        "Figure 1. End-to-end inference workflow from eligible session predictors to an actionable alert rule.",
        width=6.85,
        alt_text="Five-stage workflow showing session predictors, preprocessing, random forest, risk probability, and alert rule.",
    )

    add_page_break(doc)
    doc.add_heading("2 Validation Design", level=1)
    add_body(
        doc,
        "The split unit is the patient rather than the session. This prevents a patient with many treatments from contributing sessions to both training and evaluation partitions. A deterministic search over split seeds selected seed 30 because it produced balanced row counts and outcome prevalence across five patient-grouped folds. Fold 0 became the test set, fold 1 the validation set, and folds 2 through 4 the training set."
    )
    split_order = ["training", "validation", "test"]
    split_display = split.set_index("split").loc[split_order].reset_index()
    split_rows = [
        [
            row.split.title(),
            f"{int(row.patients):,}",
            f"{int(row.sessions):,}",
            f"{int(row.events):,}",
            f"{row.prevalence:.2%}",
        ]
        for row in split_display.itertuples(index=False)
    ]
    add_table(
        doc,
        ["Partition", "Patients", "Sessions", "Events", "Prevalence"],
        split_rows,
        widths=[1.35, 1.05, 1.35, 1.15, 1.4],
        font_size=9.0,
    )
    doc.add_heading("Leakage Controls", level=2)
    add_bullets(
        doc,
        [
            "Outcome measurements occur after the index record and are never candidate predictors.",
            "Longitudinal history features use only earlier sessions and are shifted within patient.",
            "Imputation, encoding, scaling, L1 selection, and PCA are fitted inside each training fold.",
            "Hyperparameters use training patients only; validation patients select calibration and the operating threshold.",
            "The test set remains untouched until the model, calibration method, and threshold are fixed.",
        ],
    )

    add_page_break(doc)
    doc.add_heading("3 Candidate Models", level=1)
    candidates = [
        ["Dummy prevalence baseline", "No-skill reference", "Confirms that model performance exceeds event prevalence"],
        ["Logistic regression", "Regularized linear classifier", "Transparent baseline with stable probabilities"],
        ["L1 feature selection", "Embedded sparse linear pipeline", "Tests whether a smaller feature set preserves performance"],
        ["Decision tree", "Single nonlinear tree", "Interpretable nonlinear benchmark with pruning controls"],
        ["PCA logistic regression", "Dimensionality-reduced linear model", "Tests compression after imputation and scaling"],
        ["Random forest", "Bagged tree ensemble", "Captures nonlinearities and interactions while resisting single-tree variance"],
        ["Histogram gradient boosting", "Boosted tree ensemble", "Strong tabular benchmark with efficient native implementation"],
    ]
    add_table(doc, ["Candidate", "Role", "Reason included"], candidates, widths=[2.0, 2.05, 2.85], font_size=8.2)
    doc.add_heading("Methods Not Selected", level=2)
    add_body(
        doc,
        "Kernel SVM was not prioritized because probability calibration and tuning scale poorly across more than 100,000 repeated session records, while regularized logistic regression already supplies a linear-margin benchmark. Clustering and recommender methods do not address a prespecified binary outcome. RNNs, CNNs, LSTMs, and transformers were not justified because the input is a compact tabular snapshot rather than raw images, waveforms, text, or a uniformly sampled sequence; the effective independent sample is 830 patients, and explainability is a clinical requirement. XGBoost can be evaluated later, but native histogram gradient boosting provides the same broad boosted-tree comparison with fewer external dependencies."
    )

    add_page_break(doc)
    doc.add_heading("4 Preprocessing and Tuning", level=1)
    preprocessing = [
        ["Numeric fields", "Median imputation", "Fitted independently within every fold"],
        ["Categorical fields", "Most-frequent imputation and one-hot encoding", "Unknown categories are ignored safely"],
        ["Linear candidates", "Standard scaling", "Supports regularization and comparable coefficient scales"],
        ["Tree candidates", "No scaling", "Preserves natural clinical split values"],
        ["Prior session count", "log one plus count", "Reduces strong right skew using a deterministic transformation"],
        ["PCA candidate", "Imputation scaling and 85% or 95% variance threshold", "Dimensionality reduction remains inside the pipeline"],
    ]
    add_table(doc, ["Input", "Transformation", "Reason"], preprocessing, widths=[1.55, 2.8, 2.55], font_size=8.6)
    add_body(
        doc,
        "The training partition uses three-fold StratifiedGroupKFold cross-validation. Grid search is used for the smaller linear and PCA spaces. Randomized search is used for the larger tree ensembles. All searches optimize average precision with random state 42 and two parallel workers."
    )
    tuning_rows = [
        ["Logistic regression", "C and class weight"],
        ["L1 selection", "Selector C and final-model C"],
        ["Decision tree", "Depth, leaf size, class weight, and pruning alpha"],
        ["PCA logistic", "Variance threshold, C, and class weight"],
        ["Random forest", "Depth, leaf size, feature fraction, and class weight with 250 trees"],
        ["Histogram gradient boosting", "Learning rate, leaf nodes, leaf size, L2 penalty, and class weight"],
    ]
    add_table(doc, ["Candidate", "Tuned parameters"], tuning_rows, widths=[2.25, 4.65], font_size=8.8)

    add_page_break(doc)
    doc.add_heading("5 Model Comparison", level=1)
    add_body(
        doc,
        "Average precision was the selection metric because it directly summarizes precision-recall performance under class imbalance. Brier score was the probability-quality tie-breaker. Accuracy was not used for selection because a no-skill classifier already achieves approximately 91.5% accuracy by predicting no event for every session."
    )
    add_figure(
        doc,
        ARTIFACTS / "step4_model_comparison.png",
        "Figure 2. Grouped cross-validation average precision and validation Brier score across modelling candidates.",
        width=6.85,
        alt_text="Two horizontal bar charts compare average precision and Brier score across seven candidates.",
    )
    ordered = comparison.sort_values("cv_average_precision_mean", ascending=False)
    comp_rows = [
        [
            row.model,
            f"{row.cv_average_precision_mean:.3f} +/- {row.cv_average_precision_sd:.3f}",
            f"{row.validation_average_precision:.3f}",
            f"{row.validation_roc_auc:.3f}",
            f"{row.validation_brier_score:.3f}",
        ]
        for row in ordered.itertuples(index=False)
    ]
    add_table(
        doc,
        ["Model", "Grouped CV AP", "Validation AP", "Validation AUC", "Validation Brier"],
        comp_rows,
        widths=[2.2, 1.45, 1.1, 1.1, 1.15],
        font_size=7.9,
        vertical_margin=60,
    )

    add_page_break(doc)
    doc.add_heading("6 Model Selection", level=1)
    hgb = comparison.loc[comparison.model.eq("Histogram gradient boosting")].iloc[0]
    rf = comparison.loc[comparison.model.eq("Random forest")].iloc[0]
    add_body(
        doc,
        f"Histogram gradient boosting produced the highest grouped cross-validation average precision at {hgb.cv_average_precision_mean:.3f}. Random forest reached {rf.cv_average_precision_mean:.3f}, a difference of only {hgb.cv_average_precision_mean - rf.cv_average_precision_mean:.3f}, and therefore entered the prespecified 0.01 shortlist. On the patient-disjoint validation set, random forest had higher average precision ({rf.validation_average_precision:.3f} versus {hgb.validation_average_precision:.3f}) and substantially lower Brier score ({rf.validation_brier_score:.3f} versus {hgb.validation_brier_score:.3f}). Random forest was consequently selected before any test labels were examined."
    )
    selected_params = json.loads(rf.best_parameters)
    parameter_rows = [
        ["Number of trees", "250"],
        ["Maximum depth", "Unlimited with leaf-size regularization" if selected_params.get("model__max_depth") is None else str(selected_params.get("model__max_depth"))],
        ["Minimum samples per leaf", str(selected_params["model__min_samples_leaf"])],
        ["Features considered at each split", str(selected_params["model__max_features"])],
        ["Class weighting", "None" if selected_params.get("model__class_weight") is None else str(selected_params.get("model__class_weight"))],
        ["Random state", "42"],
    ]
    add_table(doc, ["Setting", "Selected value"], parameter_rows, widths=[3.0, 3.9], font_size=9.0)
    add_body(
        doc,
        "This choice is defensible because DIAL-ALERT needs both useful risk ranking and probabilities that can support threshold or capacity decisions. It is not a claim that random forest will remain best in another centre. External validation may change the preferred model or its tuning."
    )

    add_page_break(doc)
    doc.add_heading("7 Calibration and Operating Threshold", level=1)
    add_body(
        doc,
        "Validation patients were split again by patient. One half fitted sigmoid and isotonic calibration mappings; the other half compared those mappings and selected the F2 threshold. The unmodified random-forest probabilities had the lowest decision-subset Brier score, so no post-hoc calibration was applied."
    )
    cal_rows = [
        [
            row.method.title(),
            f"{row.decision_brier_score:.4f}",
            f"{row.decision_log_loss:.4f}",
            f"{row.decision_average_precision:.3f}",
        ]
        for row in calibration.itertuples(index=False)
    ]
    add_table(doc, ["Method", "Brier", "Log loss", "Average precision"], cal_rows, widths=[1.8, 1.55, 1.55, 2.0], font_size=9.0)
    add_figure(
        doc,
        ARTIFACTS / "step4_calibration_plot.png",
        "Figure 3. Test-set calibration curve for the selected random-forest probabilities.",
        width=5.45,
        alt_text="Calibration plot compares predicted probability with observed event frequency in ten bins.",
    )
    add_body(
        doc,
        f"The threshold of {threshold:.3f} maximized F2 on the held-out validation decision subset. F2 weights recall more heavily than precision, matching the safety-oriented aim of identifying more at-risk sessions while still constraining false alerts. The threshold is a prototype operating point, not a universal clinical cutoff."
    )

    add_page_break(doc)
    doc.add_heading("8 Final Test Performance", level=1)
    add_body(
        doc,
        "The following results are from the untouched patient-disjoint test set. Confidence intervals use 500 bootstrap resamples of entire patients, preserving the dependence among repeated sessions from the same person. The intervals are wider than row-level intervals and therefore provide a more credible statement of uncertainty."
    )
    test_rows = [
        ["Average precision", fmt_ci(intervals.loc["average_precision"])],
        ["ROC AUC", fmt_ci(intervals.loc["roc_auc"])],
        ["Brier score", fmt_ci(intervals.loc["brier_score"])],
        ["Sensitivity", fmt_ci(intervals.loc["sensitivity"])],
        ["Specificity", fmt_ci(intervals.loc["specificity"])],
        ["Precision", fmt_ci(intervals.loc["precision"])],
        ["F1 score", fmt_ci(intervals.loc["f1"])],
        ["Accuracy", f"{metrics['accuracy']:.3f}"],
    ]
    add_table(doc, ["Metric", "Estimate and 95% interval"], test_rows, widths=[3.0, 3.9], font_size=9.0)
    add_figure(
        doc,
        ARTIFACTS / "test_performance_intervals.png",
        "Figure 4. Final performance estimates with patient-cluster bootstrap 95% confidence intervals.",
        width=5.35,
        alt_text="Error-bar chart shows test metrics and 95 percent patient-cluster bootstrap intervals.",
    )

    add_page_break(doc)
    doc.add_heading("9 Discrimination and Classification", level=1)
    add_figure(
        doc,
        ARTIFACTS / "roc_pr_curves.png",
        "Figure 5. ROC and precision-recall curves on the untouched test set.",
        width=6.7,
        alt_text="Side-by-side ROC and precision-recall curves for the selected random forest.",
    )
    add_body(
        doc,
        f"ROC AUC of {metrics['roc_auc']:.3f} indicates good rank discrimination. Average precision of {metrics['average_precision']:.3f} is more informative for this imbalanced outcome and is well above the test prevalence of {int(split.loc[split.split.eq('test'), 'events'].iloc[0]) / int(split.loc[split.split.eq('test'), 'sessions'].iloc[0]):.3f}. The lower test average precision than validation average precision is a reminder that performance varies across patient groups and should not be summarized by a single optimistic split."
    )
    add_figure(
        doc,
        ARTIFACTS / "confusion_matrix.png",
        f"Figure 6. Confusion matrix at the validation-selected threshold of {threshold:.3f}.",
        width=5.2,
        alt_text=f"Confusion matrix with {metrics['tn']} true negatives, {metrics['fp']} false positives, {metrics['fn']} false negatives, and {metrics['tp']} true positives.",
    )
    add_body(
        doc,
        f"At this threshold, the model identified {metrics['tp']:,} of {metrics['tp'] + metrics['fn']:,} events and produced {metrics['fp']:,} false alerts among {metrics['tn'] + metrics['fp']:,} non-event sessions. This trade-off reflects the F2 objective. A dialysis unit may choose another threshold after defining acceptable missed-event and alert-burden costs."
    )

    doc.add_heading("10 Alert Capacity Analysis", level=1)
    add_body(
        doc,
        "Capacity analysis ranks sessions by risk and asks what happens if staff can review only a fixed fraction. This avoids implying that one threshold fits every unit and converts model performance into workload and captured-event terms."
    )
    add_figure(
        doc,
        ARTIFACTS / "capacity_curve.png",
        "Figure 7. Event recall and positive predictive value as alert capacity increases.",
        width=6.15,
        alt_text="Line chart shows recall increasing and precision decreasing as alert capacity rises from 5 to 50 percent.",
    )
    selected_caps = capacity[capacity.alert_capacity.isin([0.05, 0.10, 0.20, 0.30, 0.50])]
    cap_rows = [
        [
            f"{row.alert_capacity:.0%}",
            f"{int(row.selected_sessions):,}",
            f"{row.precision_at_capacity:.1%}",
            f"{row.recall_at_capacity:.1%}",
            f"{row.lift_at_capacity:.2f}",
            f"{row.false_alerts_per_100_sessions:.1f}",
        ]
        for row in selected_caps.itertuples(index=False)
    ]
    add_table(
        doc,
        ["Capacity", "Alerts", "Precision", "Recall", "Lift", "False alerts per 100"],
        cap_rows,
        widths=[0.85, 0.95, 1.05, 1.0, 0.8, 1.6],
        font_size=8.3,
    )
    add_body(
        doc,
        f"At 20% capacity, {int(metrics['selected_sessions']):,} test sessions receive an alert. The strategy captures {metrics['recall_at_capacity']:.1%} of events with {metrics['precision_at_capacity']:.1%} precision and {metrics['lift_at_capacity']:.2f}-fold lift. It generates approximately {metrics['false_alerts_per_100_sessions']:.1f} false alerts per 100 total sessions, which must be assessed against staffing and intervention burden."
    )

    add_page_break(doc)
    doc.add_heading("11 Model Interpretation", level=1)
    add_body(
        doc,
        "Permutation importance was calculated on a reproducible test-set sample using average precision as the score. Previous-session nadir SBP and prior IDH rate were the strongest unique contributors, followed by fluid excess percentage and baseline haemodynamic measures. This agrees with the Step 3 EDA and supports face validity, but it does not establish causal effects."
    )
    add_figure(
        doc,
        ARTIFACTS / "permutation_importance.png",
        "Figure 8. Permutation importance for the selected random-forest pipeline.",
        width=5.95,
        alt_text="Horizontal bars show prior nadir SBP and prior IDH rate as the most important predictors.",
    )
    imp_rows = [
        [row.feature.replace("_", " ").title(), f"{row.importance_mean:.3f}", f"{row.importance_sd:.3f}"]
        for row in importance.head(6).itertuples(index=False)
    ]
    add_table(doc, ["Feature", "Mean decrease in AP", "Standard deviation"], imp_rows, widths=[3.6, 1.7, 1.6], font_size=8.8)

    add_page_break(doc)
    doc.add_heading("12 Reproducibility and Saved Artefacts", level=1)
    add_body(
        doc,
        f"The run used Python {manifest['python_version']}, scikit-learn {manifest['package_versions']['scikit_learn']}, random state {manifest['random_state']}, split seed {manifest['split_seed']}, and {manifest['grouped_cv_folds']}-fold grouped cross-validation. The manifest records SHA-256 hashes for the processed data, configuration, and model files so accidental changes can be detected."
    )
    files_rows = [
        ["models/dial_alert_final_predictor.joblib", "Selected fitted preprocessing and random-forest pipeline"],
        ["models/decision_threshold.json", "F2 threshold and ordered feature list"],
        ["models/model_manifest.json", "Versions, seeds, sizes, and SHA-256 hashes"],
        ["artifacts/model_comparison.csv", "Grouped-CV and validation metrics for all candidates"],
        ["artifacts/cv_search_top_results.csv", "Top hyperparameter-search results"],
        ["artifacts/final_test_metrics.json", "Locked test-set results"],
        ["artifacts/test_metric_confidence_intervals.csv", "Patient-cluster bootstrap intervals"],
        ["artifacts/capacity_metrics.csv", "Alert workload and captured-event trade-offs"],
        ["artifacts/split_assignments.csv.gz", "Reproducible patient-disjoint assignment for every session"],
        ["configs/model_config.json", "Features, target, capacity, seeds, workers, folds, and selection tolerance"],
    ]
    add_table(doc, ["Artefact", "Purpose"], files_rows, widths=[2.7, 4.2], font_size=8.0, vertical_margin=60)
    doc.add_heading("Execution Commands", level=2)
    add_code(
        doc,
        [
            "python src/build_session_dataset.py --raw-dir data/raw --output-dir data/processed",
            "python src/train_evaluate.py --data data/processed/hemobp_session_level.csv.gz --config configs/model_config.json --artifacts artifacts --models models",
            "python src/generate_step4_assets.py",
            "python src/predict.py --input new_sessions.csv --output predictions.csv",
        ],
    )

    add_page_break(doc)
    doc.add_heading("13 Limitations", level=1)
    limitation_rows = [
        ["Retrospective single-centre data", "Performance may not transport to other dialysis units, eras, or patient populations.", "External temporal and geographic validation"],
        ["BP-defined target", "Symptoms and treatments required for a complete clinical syndrome are unavailable.", "Prospective data collection with symptom and intervention fields"],
        ["Informative measurement", "Unstable sessions may have more BP readings and more opportunities to detect events.", "Fixed-window sensitivity analyses and prospective protocolized measurement"],
        ["Repeated sessions", "A large session count does not equal the same number of independent patients.", "Patient-grouped validation and cluster-bootstrap intervals"],
        ["Cold start", "Prior-session features are missing for first observed sessions.", "Report cold-start performance and consider a reduced-input model"],
        ["Threshold dependence", "Sensitivity and false-alert burden change with workflow capacity and costs.", "Site-specific utility study before threshold adoption"],
        ["No causal interpretation", "Importance and partial effects describe prediction, not treatment benefit.", "Prospective impact evaluation before clinical claims"],
    ]
    add_table(doc, ["Limitation", "Effect", "Required response"], limitation_rows, widths=[1.4, 2.85, 2.65], font_size=7.8)
    add_body(
        doc,
        "The study demonstrates model development and internal validation. It does not demonstrate that acting on an alert improves patient outcomes, reduces costs, or is safe in routine care."
    )

    add_page_break(doc)
    doc.add_heading("14 Conclusions", level=1)
    add_body(
        doc,
        f"DIAL-ALERT completed the supervised model-implementation stage with a reproducible comparison of seven configurations. {selected} was chosen by a locked rule that considered both grouped cross-validation ranking performance and validation probability error. On patient-disjoint testing, the model achieved ROC AUC {metrics['roc_auc']:.3f}, average precision {metrics['average_precision']:.3f}, and Brier score {metrics['brier_score']:.3f}."
    )
    add_body(
        doc,
        f"The validation-selected threshold detected {metrics['sensitivity']:.1%} of events at {metrics['precision']:.1%} precision. A capacity-based policy that alerts on the highest-risk 20% captured {metrics['recall_at_capacity']:.1%} of events with {metrics['lift_at_capacity']:.2f}-fold lift. These are promising internal results, but the uncertainty intervals and false-alert burden require careful interpretation."
    )
    add_body(
        doc,
        "The trained pipelines, threshold, configuration, split assignments, metrics, package versions, and integrity hashes are saved. The next capstone stage should audit model explanations, limitations, subgroup fairness, and feasible mitigation strategies before any deployment work begins."
    )

    doc.add_heading("References", level=1)
    references = [
        "1. Lin CJ, Chen YY, Pan CF, Wu VC, Wu CJ. Dataset supporting blood pressure prediction for the management of chronic hemodialysis. Scientific Data. 2019;6:313. https://doi.org/10.1038/s41597-019-0319-8",
        "2. Breiman L. Random forests. Machine Learning. 2001;45:5-32. https://doi.org/10.1023/A:1010933404324",
        "3. Saito T, Rehmsmeier M. The precision-recall plot is more informative than the ROC plot when evaluating binary classifiers on imbalanced datasets. PLoS ONE. 2015;10:e0118432. https://doi.org/10.1371/journal.pone.0118432",
        "4. Brier GW. Verification of forecasts expressed in terms of probability. Monthly Weather Review. 1950;78:1-3.",
        "5. Varma S, Simon R. Bias in error estimation when using cross-validation for model selection. BMC Bioinformatics. 2006;7:91. https://doi.org/10.1186/1471-2105-7-91",
        "6. Pedregosa F and colleagues. Scikit-learn machine learning in Python. Journal of Machine Learning Research. 2011;12:2825-2830.",
    ]
    for reference in references:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(-14)
        p.paragraph_format.left_indent = Pt(14)
        set_font(p.add_run(reference), size=8.8)

    doc.core_properties.title = "DIAL ALERT Model Implementation Report"
    doc.core_properties.subject = "Capstone Project Step 4"
    doc.core_properties.author = "Franklin Guillano"
    doc.core_properties.keywords = "DIAL-ALERT, HEMOBP, random forest, haemodialysis, model comparison"
    doc.core_properties.comments = "DIAL-ALERT academic capstone"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(make_report())
