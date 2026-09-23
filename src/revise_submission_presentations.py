"""Update report wording and managed notes without retraining.

Run from the project checkout:
    python src/revise_submission_documents.py --check
    python src/revise_submission_documents.py

Existing notes are updated in place; missing notes are added only once.
Only changed DOCX files are saved. Originals are backed up outside the repository
in ~/DIAL_ALERT_document_backups. PDF files are not changed by this script.
"""
from __future__ import annotations

import argparse
import os
import shutil
import tempfile
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph
from docx.text.run import Run

ROOT = Path(__file__).resolve().parents[1]
REPLACEMENTS = {
    "a twelve-slide technical deck for peers and a ten-slide business deck": "a sixteen-slide technical deck for peers and an eleven-slide business deck",
    "0.444 for histogram gradient boosting": "0.446 for histogram gradient boosting",
    "Previous-session nadir SBP produced the largest decrease (0.128), followed by prior IDH rate (0.081), fluid excess percentage (0.015), baseline SBP (0.009), and baseline mean arterial pressure (0.004).": "Previous-session nadir SBP produced the largest decrease (0.076), followed by prior IDH rate (0.059), previous-session IDH (0.017), baseline SBP (0.014), and baseline mean arterial pressure (0.009).",
    "Configurations, split assignments, candidate pipelines, final predictor, threshold, package versions, hashes, and random seeds are saved.": "The final predictor, threshold, configuration and version records are included. Candidate pipelines and patient-level split assignments must be regenerated.",
    "Reproducible patient-disjoint assignment for every session": "Regenerated locally; patient-level assignments are not bundled",
    "The trained pipelines, threshold, configuration, split assignments, metrics, package versions, and integrity hashes are saved.": "The selected predictor, threshold, configuration, metrics and version records are included; other pipelines and patient-level split assignments require regeneration.",
    "follow-up shorter than 120 minutes": "last observed dialysis minute below 120",
    "observation through minute 120": "an observation at or beyond dialysis minute 120 (not 120 minutes after prediction)",
    "observation through at least minute 120": "an observation at or beyond dialysis minute 120",
    "follow-up to at least minute 120": "an observation at or beyond dialysis minute 120 (not 120 minutes after prediction)",
    "Sessions removed for follow-up below 120 minutes": "Sessions with last observed dialysis minute below 120",
    "This rule balanced event ranking with probability reliability and selected the random forest over the less well-calibrated boosted-tree candidate.": "This rule selected random forest using the candidates as fitted. It does not establish superiority after equal calibration of both finalists.",
    "SHAP explanations were generated for the uncalibrated base model because probability calibration does not change the underlying predictor relationships.": "SHAP explanations describe the selected random forest; no post-hoc calibration was applied.",
    "Appendix Rubric Evidence Map": "Appendix Project Evidence Map",
    "Rubric step": "Project step",
}

NOTES = [
    ("Outcome and eligibility", "Predict any SBP below 90 mmHg after the earliest valid active-dialysis observation in minutes 0 to 30. Require index SBP at least 90, two distinct later measurement minutes, and last observed dialysis minute at least 120. The last requirement is anchored to dialysis time, not to prediction time. It is a retrospective eligibility criterion; short or interrupted sessions need separate prospective evaluation."),
    ("Distinct operating strategies", "The fixed threshold is 0.142855878 (rounded to 0.143), chosen by maximum F2 on a patient-disjoint validation decision subset: test sensitivity 66.1%, precision 30.6%, 18.3 reviews and 12.7 false alerts per 100 sessions. Reviewing the highest-risk 20% instead gives 69.1% recall, 29.4% precision and 14.1 false alerts per 100. A prospective capacity policy must define the comparison batch and tie handling before use."),
    ("Model choice and uncertainty", "Random forest was selected from candidates within 0.01 of the best grouped-CV average precision using validation Brier score. No post-hoc calibration was retained. An equally calibrated finalist comparison remains secondary work. The 21,354 test sessions come from only 170 patients. Patient-bootstrap 95% intervals: AP 0.306 to 0.478; ROC AUC 0.821 to 0.876; sensitivity at the fixed threshold 55.0% to 74.5%; recall at 20% capacity 63.2% to 74.4%."),
    ("Proposed future pilot targets", "Provisional targets, not demonstrated prospective results: detect at least 65% of later SBP-below-90 events, generate no more than 15 false alerts per 100 eligible sessions, and achieve median review time at most 2 minutes per displayed alert. Prespecify one operating policy, denominators, uncertainty and safety stopping rules with the local team before evaluation."),
]
FUTURE = [
    ("Optional deployment evidence", "Step 8 includes a local Flask app, a synthetic request, a recorded HTTP demo and instructions. The app and command-line route apply the same prior-session-count transformation. This is a local inference demonstration; no clinical or cloud deployment is established. Step 9 includes an earlier saved-draft replay in step9_genai and a later live local Ollama assistant demonstration in step9_assistant. The live demonstration includes examples, review evidence, and an edited video. Observed answer-quality limitations remain; citation and numeric checks do not establish semantic correctness."),
    ("Reproducibility status", "A fresh end-to-end reproduction was successfully completed using Python 3.12 on 20 September 2026. The HEMOBP Version 3 source files were downloaded and checksum-verified, the session-level analytical dataset was rebuilt from the raw data, models were retrained and evaluated, EDA and fairness outputs were regenerated, and the automated test suite completed with 15 passed, 0 failed, and 0 skipped. The fresh run was numerically consistent with the locked reference results rather than byte-for-byte identical. Full details and numerical comparisons are recorded in docs/REPRODUCIBILITY_RECORD.md."),
    ("Secondary analyses awaiting execution", "Measure index-to-first-event warning time; compare a baseline-SBP plus prior-hypotension model; assess history ablation and first-observed sessions; compare both finalists with identical calibration splits and methods; and examine extreme UF/fluid-excess values and short-session exclusion. Prespecify analyses on development data, report them as exploratory, and do not replace the locked model based on repeated test-set comparisons."),
    ("Future validation sequence", "Validate the locked model at another center or in a later period, confirm exploratory fairness mitigation on fresh patients, then run silent predictions. Only after data quality and safety gates pass should supervised clinician review assess workload, actions, unintended effects, outcomes and costs. A 90-day schedule is illustrative and cannot establish clinical or financial benefit by itself."),
]
SECTIONS = [
    ("Submission interpretation and evidence", NOTES),
    ("Demonstration status and further research", FUTURE),
]


class RevisionError(ValueError):
    """A report needs manual inspection rather than a potentially destructive edit."""


def paragraphs(doc: DocumentObject):
    """Yield body and nested-table paragraphs once, including merged cells once.

    Header/footer text, text boxes, comments and tracked changes are not edited.
    """
    seen = set()

    def walk(container):
        for paragraph in container.paragraphs:
            if paragraph._p not in seen:
                seen.add(paragraph._p)
                yield paragraph
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from walk(cell)

    yield from walk(doc)


def _visible_runs(p: Paragraph) -> list[Run]:
    # Include hyperlink runs, which p.runs alone does not expose.
    runs = [Run(element, p) for element in p._p.xpath('./w:r | ./w:hyperlink/w:r')]
    if ''.join(run.text for run in runs) != p.text:
        raise RevisionError('Unsupported text structure in a paragraph needing correction.')
    return runs


def _assert_plain_run(run: Run) -> None:
    allowed = {qn(name) for name in ('w:rPr', 'w:t', 'w:tab', 'w:br', 'w:cr')}
    if any(child.tag not in allowed for child in run._r):
        raise RevisionError('A correction touches a field or embedded object; inspect it manually.')


def _replace_span(p: Paragraph, start: int, end: int, new: str) -> None:
    """Change a character span, keeping run formatting outside the changed span."""
    runs = _visible_runs(p)
    if not (0 <= start <= end <= len(p.text)):
        raise RevisionError('Invalid replacement span.')
    if start == end:
        offset = 0
        for run in runs:
            if offset <= start < offset + len(run.text):
                _assert_plain_run(run)
                at = start - offset
                run.text = run.text[:at] + new + run.text[at:]
                return
            offset += len(run.text)
        # End-of-paragraph insertion must not extend an existing hyperlink.
        p.add_run(new)
        return
    offset = 0
    for run in runs:
        text = run.text
        left, right = offset, offset + len(text)
        offset = right
        if right <= start or left >= end:
            continue
        _assert_plain_run(run)
        prefix = text[:max(0, start - left)]
        suffix = text[max(0, end - left):] if right > end else ''
        run.text = prefix + (new if left <= start < right else '') + suffix


def replace_text(p: Paragraph, old: str, new: str) -> int:
    """Replace original matches once, right-to-left; never recurse indefinitely."""
    if not old:
        raise ValueError('The text to replace must not be empty.')
    if old == new or old not in p.text:
        return 0
    original = p.text
    starts = []
    cursor = 0
    while (start := original.find(old, cursor)) != -1:
        starts.append(start)
        cursor = start + len(old)
    for start in reversed(starts):
        _replace_span(p, start, start + len(old), new)
    if p.text != original.replace(old, new):
        raise RevisionError('Text replacement did not produce the expected result.')
    return len(starts)


def _set_body(p: Paragraph, body: str) -> None:
    """Keep unchanged prefix/suffix formatting when refreshing a managed note."""
    old = p.text
    start = 0
    while start < min(len(old), len(body)) and old[start] == body[start]:
        start += 1
    tail = 0
    while (tail < min(len(old), len(body)) - start
           and old[-tail - 1] == body[-tail - 1]):
        tail += 1
    old_end = len(old) - tail
    new_end = len(body) - tail
    _replace_span(p, start, old_end, body[start:new_end])
    if p.text != body:
        raise RevisionError('The managed note was not updated completely.')


def _level(p: Paragraph) -> int | None:
    style = p.style
    while style is not None:
        name = style.name or ''
        if name.startswith('Heading ') and name[8:].isdigit():
            return int(name[8:])
        style = style.base_style
    return None


def _section(doc: DocumentObject, title: str):
    matches = [p for p in doc.paragraphs if p.text.strip() == title]
    if len(matches) != 1:
        raise RevisionError(f'Expected one section heading {title!r}; found {len(matches)}.')
    heading = matches[0]
    if _level(heading) != 1:
        raise RevisionError(f'{title!r} is not a Heading 1 paragraph; inspect its structure.')
    contents = []
    for element in heading._p.itersiblings():
        if element.tag == qn('w:sectPr'):
            break
        if element.tag == qn('w:p'):
            paragraph = Paragraph(element, doc._body)
            if _level(paragraph) == 1:
                break
        contents.append(element)
    return heading, contents


def _append_or_insert(doc: DocumentObject, text: str, style: str,
                      before=None) -> Paragraph:
    paragraph = doc.add_paragraph(text, style=style)
    if before is not None:
        before.addprevious(paragraph._p)
    return paragraph


def _before_break(element):
    """Keep a section's pre-existing page-break paragraph with that section."""
    if element is not None:
        previous = element.getprevious()
        if (previous is not None and previous.tag == qn('w:p')
                and previous.xpath('.//w:br[@w:type="page"]')
                and not previous.xpath('.//w:t | .//w:drawing | .//w:pict')):
            return previous
    return element


def _section_end(doc: DocumentObject, title: str):
    heading, contents = _section(doc, title)
    last = contents[-1] if contents else heading._p
    following = last.getnext()
    # Move the insertion point before a trailing page break, not after it.
    if following is not None:
        return _before_break(following)
    return None


def _note_body(doc: DocumentObject, title: str, heading_text: str):
    _, contents = _section(doc, title)
    matches = [element for element in contents
               if element.tag == qn('w:p')
               and Paragraph(element, doc._body).text.strip() == heading_text]
    if len(matches) != 1:
        raise RevisionError(f'Expected one note heading {heading_text!r}; found {len(matches)}.')
    heading = Paragraph(matches[0], doc._body)
    if _level(heading) != 2:
        raise RevisionError(f'{heading_text!r} is not a Heading 2 paragraph.')
    bodies = []
    for element in contents[contents.index(heading._p) + 1:]:
        if element.tag != qn('w:p'):
            raise RevisionError(f'Unexpected table/object in note {heading_text!r}; inspect manually.')
        paragraph = Paragraph(element, doc._body)
        if _level(paragraph) is not None:
            break
        if paragraph.text.strip():
            bodies.append(paragraph)
        elif element.xpath('.//w:drawing | .//w:pict | .//w:object'):
            raise RevisionError(f'Unexpected image/object in note {heading_text!r}.')
    if len(bodies) > 1:
        raise RevisionError(f'Note {heading_text!r} has multiple body paragraphs; no text was deleted.')
    return heading, bodies[0] if bodies else None


def _format_new_body(p: Paragraph) -> None:
    p.paragraph_format.space_after = Pt(7)
    for run in p.runs:
        run.font.size = Pt(10)


def append_notes(doc: DocumentObject) -> dict[str, int]:
    """Compatibility name: synchronize existing notes rather than blindly append."""
    counts = {'sections_added': 0, 'notes_added': 0, 'notes_updated': 0}
    for section_index, (title, rows) in enumerate(SECTIONS):
        matches = [p for p in doc.paragraphs if p.text.strip() == title]
        if len(matches) > 1:
            raise RevisionError(f'Duplicate section {title!r}; inspect it before saving.')
        if not matches:
            later_titles = {name for name, _ in SECTIONS[section_index + 1:]}
            following = next((p._p for p in doc.paragraphs
                              if p.text.strip() in later_titles), None)
            section_heading = _append_or_insert(doc, title, 'Heading 1', _before_break(following))
            section_heading.paragraph_format.page_break_before = True
            counts['sections_added'] += 1
        for note_index, (heading_text, body) in enumerate(rows):
            _, contents = _section(doc, title)
            matches = [element for element in contents
                       if element.tag == qn('w:p')
                       and Paragraph(element, doc._body).text.strip() == heading_text]
            if len(matches) > 1:
                raise RevisionError(f'Duplicate note {heading_text!r}; inspect it before saving.')
            if not matches:
                later_names = {name for name, _ in rows[note_index + 1:]}
                following = next((element for element in contents
                                  if element.tag == qn('w:p')
                                  and Paragraph(element, doc._body).text.strip() in later_names), None)
                if following is None:
                    following = _section_end(doc, title)
                new_heading = _append_or_insert(doc, heading_text, 'Heading 2', _before_break(following))
                paragraph = doc.add_paragraph(body)
                new_heading._p.addnext(paragraph._p)
                _format_new_body(paragraph)
                counts['notes_added'] += 1
                continue
            heading, paragraph = _note_body(doc, title, heading_text)
            if paragraph is None:
                paragraph = doc.add_paragraph(body)
                heading._p.addnext(paragraph._p)
                _format_new_body(paragraph)
                counts['notes_added'] += 1
            elif paragraph.text != body:
                _set_body(paragraph, body)
                counts['notes_updated'] += 1
    validate_notes(doc)
    return counts


def validate_notes(doc: DocumentObject) -> None:
    for title, rows in SECTIONS:
        for heading, body in rows:
            _, paragraph = _note_body(doc, title, heading)
            if paragraph is None or paragraph.text != body:
                raise RevisionError(f'Missing or outdated note: {heading!r}.')


def revise(doc: DocumentObject, final_report: bool) -> dict[str, int]:
    counts = {'text_replacements': 0, 'sections_added': 0, 'notes_added': 0, 'notes_updated': 0}
    for paragraph in paragraphs(doc):
        for old, new in REPLACEMENTS.items():
            counts['text_replacements'] += replace_text(paragraph, old, new)
    if final_report:
        counts.update(append_notes(doc))
    return counts


def _serialize_document(original: bytes, doc: DocumentObject) -> bytes:
    # Edits only change body XML and use existing styles/relationships. Preserve
    # every other ZIP member verbatim instead of reserializing the whole package.
    stream = BytesIO()
    part_name = str(doc.part.partname).lstrip('/')
    with ZipFile(BytesIO(original)) as source, ZipFile(stream, 'w') as target:
        if any(name.startswith('_xmlsignatures/') for name in source.namelist()):
            raise RevisionError('Digitally signed DOCX: revise and re-sign it manually.')
        target.comment = source.comment
        for info in source.infolist():
            payload = doc.part.blob if info.filename == part_name else source.read(info.filename)
            target.writestr(info, payload)
    return stream.getvalue()


def prepare(path: Path):
    original = path.read_bytes()
    doc = Document(BytesIO(original))
    final_report = 'Final_Report' in path.name
    counts = revise(doc, final_report)
    if not any(counts.values()):
        return path, original, original, counts
    data = _serialize_document(original, doc)
    # Reopen and ensure that another run would make no further edits.
    checked = Document(BytesIO(data))
    before = checked.element.xml
    if any(revise(checked, final_report).values()) or checked.element.xml != before:
        raise RevisionError(f'Repeat-run stability check failed for {path.name}.')
    return path, original, data, counts


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--check', action='store_true',
                        help='Validate proposed edits without saving reports or creating backups.')
    parser.add_argument('--reports-dir', type=Path, default=ROOT / 'reports',
                        help='Reports directory; defaults to the repository reports folder.')
    args = parser.parse_args()
    reports_dir = args.reports_dir.expanduser().resolve()
    files = sorted(path for path in reports_dir.glob('*.docx')
                   if path.is_file() and not path.name.startswith('~$'))
    if not files:
        raise FileNotFoundError(f'No DOCX reports found in {reports_dir}.')
    prepared = []
    for path in files:
        try:
            prepared.append(prepare(path))
        except Exception as exc:
            raise RevisionError(f'{path.name}: {exc}. No report files have been saved.') from exc
    changed = []
    for path, original, data, counts in prepared:
        if any(counts.values()):
            changed.append((path, original, data))
            details = ', '.join(f'{key}={value}' for key, value in counts.items() if value)
            print(f'{path.name}: {details}')
        else:
            print(f'{path.name}: unchanged')
    if args.check:
        print('CHECK ONLY: no report or backup files written.')
        return
    if not changed:
        print('All reports are already current; no files written.')
        return
    for path, original, _ in changed:
        if path.read_bytes() != original:
            raise RevisionError(f'{path.name} changed during validation; rerun after saving your edits.')
    backup_root = Path.home() / 'DIAL_ALERT_document_backups'
    backup_root.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ_')
    backup_dir = Path(tempfile.mkdtemp(prefix=stamp, dir=backup_root))
    for path, original, _ in changed:
        (backup_dir / path.name).write_bytes(original)
    print(f'Original reports backed up to: {backup_dir}')
    for path, _, data in changed:
        temporary = None
        try:
            with tempfile.NamedTemporaryFile(dir=path.parent, suffix='.tmp', delete=False) as stream:
                temporary = Path(stream.name)
                stream.write(data)
                stream.flush()
                os.fsync(stream.fileno())
            shutil.copymode(path, temporary)
            temporary.replace(path)
        finally:
            if temporary is not None and temporary.exists():
                temporary.unlink()
    print(f'Saved {len(changed)} DOCX report(s). PDFs were not changed; review exports before publishing.')


if __name__ == '__main__':
    main()
