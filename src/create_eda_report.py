"""Create the DIAL-ALERT EDA and Feature Engineering Report."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS = ROOT / "artifacts"
DOCS = ROOT / "docs"
REPORTS = ROOT / "reports"
REPORTS.mkdir(exist_ok=True)

OUTPUT = REPORTS / "Franklin_Guillano_DIAL_ALERT_EDA_Feature_Engineering_Report.docx"

NAVY = "0B5C8E"
PALE_BLUE = "EAF3F8"
PALE_GRAY = "F5F6F7"
MID_GRAY = "6B7280"
LIGHT_GRAY = "D9D9D9"
BLACK = "000000"
WHITE = "FFFFFF"


def set_font(run, name: str = "Arial", size: float | None = None, bold: bool | None = None,
             italic: bool | None = None, color: str | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 90, start: int = 100, bottom: int = 90, end: int = 100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color: str = LIGHT_GRAY, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_repeat_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_font(run, size=8.5, color=MID_GRAY)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(14)
    title_ppr = title._element.get_or_add_pPr()
    for border in list(title_ppr.findall(qn("w:pBdr"))):
        title_ppr.remove(border)

    subtitle = doc.styles["Subtitle"]
    subtitle.font.name = "Arial"
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    subtitle.font.size = Pt(15)
    subtitle.font.color.rgb = RGBColor.from_string(MID_GRAY)
    subtitle.paragraph_format.space_after = Pt(12)

    for style_name, size, before, after in (
        ("Heading 1", 18, 8, 8),
        ("Heading 2", 13.5, 8, 5),
        ("Heading 3", 11.5, 6, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    caption.font.name = "Arial"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    caption.font.size = Pt(8.5)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MID_GRAY)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("DIAL ALERT   |   EDA and Feature Engineering Report   |   ")
    set_font(r, size=8.5, color=MID_GRAY)
    set_repeat_page_number(p)


def add_body(doc: Document, text: str, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_font(r2)
    else:
        r = p.add_run(text)
        set_font(r)


def add_bullets(doc: Document, items: list[str], level: int = 0) -> None:
    for item in items:
        style = "List Bullet" if level == 0 else "List Bullet 2"
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_font(r)


def add_table(doc: Document, headers: list[str], rows: list[list[object]], widths: list[float] | None = None,
              font_size: float = 9.0, vertical_margin: int = 90) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header = table.rows[0]
    repeat_table_header(header)
    for i, value in enumerate(headers):
        cell = header.cells[i]
        cell.text = str(value)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell, top=vertical_margin, bottom=vertical_margin)
        if widths:
            cell.width = Inches(widths[i])
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_font(run, size=font_size, bold=True, color=WHITE)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for i, value in enumerate(values):
            cell = cells[i]
            cell.text = str(value)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_shading(cell, WHITE if row_index % 2 == 0 else PALE_BLUE)
            set_cell_border(cell)
            set_cell_margins(cell, top=vertical_margin, bottom=vertical_margin)
            if widths:
                cell.width = Inches(widths[i])
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 or len(str(value)) > 18 else WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    set_font(run, size=font_size)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)


def add_figure(doc: Document, path: Path, caption: str, width: float = 6.75, alt_text: str | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(path), width=Inches(width))
    if alt_text:
        shape._inline.docPr.set("descr", alt_text)
    cp = doc.add_paragraph(style="Caption")
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(7)
    cp.paragraph_format.keep_together = True
    r = cp.add_run(caption)
    set_font(r, size=8.5, italic=True, color=MID_GRAY)


def add_page_break(doc: Document) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_code(doc: Document, lines: list[str]) -> None:
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(line)
        set_font(r, name="Courier New", size=9, color=BLACK)


def make_report() -> Path:
    df = pd.read_csv(ROOT / "data/processed/hemobp_session_level.csv.gz", parse_dates=["session_date"])
    flow = json.loads((ROOT / "data/processed/cohort_flow.json").read_text(encoding="utf-8"))
    cleaning = pd.read_csv(ARTIFACTS / "cleaning_audit.csv")
    subgroup = pd.read_csv(ARTIFACTS / "subgroup_prevalence.csv")
    selected = pd.read_csv(ARTIFACTS / "l1_selected_features.csv")
    selected_names = selected.loc[selected.selected.eq(1), "feature"].tolist()
    pca = pd.read_csv(ARTIFACTS / "pca_explained_variance.csv")
    importance = pd.read_csv(ARTIFACTS / "permutation_importance.csv")
    feature_engineering = pd.read_csv(DOCS / "feature_engineering_dictionary.csv")

    doc = Document()
    configure_document(doc)

    # Cover
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(95)
    p.add_run("DIAL ALERT EDA and Feature Engineering Report")
    p2 = doc.add_paragraph(style="Subtitle")
    p2.add_run("Capstone Project Step 3")
    p3 = doc.add_paragraph()
    p3.paragraph_format.space_before = Pt(18)
    r = p3.add_run("Explainable early prediction of blood pressure defined intradialytic hypotension")
    set_font(r, size=12, bold=True)
    p4 = doc.add_paragraph()
    r = p4.add_run("Prepared by Franklin Guillano")
    set_font(r, size=11)
    p5 = doc.add_paragraph()
    r = p5.add_run("September 2026")
    set_font(r, size=10.5, color=MID_GRAY)
    p6 = doc.add_paragraph()
    p6.paragraph_format.space_before = Pt(110)
    r = p6.add_run(
        "Academic clinical decision support prototype. The output is not a diagnostic device and requires external and prospective validation before clinical use."
    )
    set_font(r, size=9.5, italic=True, color=MID_GRAY)

    add_page_break(doc)
    doc.add_heading("Executive Summary", level=1)
    add_body(
        doc,
        "This report documents the reproducible transformation of the HEMOBP source files into a leakage-controlled session-level modelling table. The final analytic cohort contains 106,758 haemodialysis sessions from 830 patients. The primary blood-pressure-defined outcome occurred in 9,067 sessions, giving a prevalence of 8.49%. Missingness was limited to dialysis vintage and first-session history variables. Physiologically impossible monitoring records were removed before feature construction, while plausible extreme values were retained and described rather than deleted automatically."
    )
    add_body(
        doc,
        "Exploratory analysis identified strong longitudinal and treatment-related associations. A lower previous-session nadir, a higher previous IDH rate, greater fluid excess, and higher initial ultrafiltration intensity were associated with greater observed event prevalence. These relationships informed domain-derived features but are not interpreted as causal effects."
    )
    add_body(
        doc,
        "The preprocessing pipeline applies training-fold median imputation, categorical mode imputation, one-hot encoding, and model-specific scaling. Embedded L1 selection retained 12 of 24 transformed features. PCA retained 11 numeric components to exceed the selected 85% variance threshold, but the PCA candidate underperformed the interpretable non-PCA alternatives. Global permutation importance, aggregate SHAP values, and dependence curves based on synthetic reference profiles provide complementary explanations of the selected model."
    )

    add_page_break(doc)
    doc.add_heading("Rubric Evidence", level=2)
    rubric_rows = [
        ["Clean data", "Nulls, dates, duplicates, linkage, physiologic bounds, index-time rules, and follow-up eligibility are documented."],
        ["Engineer features", "Scaling, encoding, diagnostic binning, longitudinal history, haemodynamic, and dialysis-dose features are reproduced in code."],
        ["Applied EDA", "Outcome distributions, predictor distributions, correlations, subgroup prevalence, nonlinear decile relationships, and index timing are shown."],
        ["Feature importance", "Permutation importance, aggregate SHAP, a synthetic local explanation, and synthetic-profile dependence analyses are included."],
        ["Feature selection", "Embedded L1 logistic regression selects 12 of 24 transformed predictors inside the training workflow."],
        ["Dimensionality reduction", "PCA is fitted after imputation and scaling; 11 components explain 87.6% of numeric-feature variance."],
        ["Reproducibility", "Source scripts, configuration, random seed, generated tables, and commands are listed."],
    ]
    add_table(doc, ["Requirement", "Evidence in this report"], rubric_rows, widths=[1.65, 5.25], font_size=8.8)

    add_page_break(doc)
    doc.add_heading("1 Analysis Objective", level=1)
    add_body(
        doc,
        "DIAL-ALERT is a supervised binary-classification project. Each eligible haemodialysis session receives a risk probability for a later systolic blood-pressure reading below 90 mmHg. The index observation is the earliest valid active-dialysis record within the first 30 minutes, and the outcome is derived only from later measurements. Sessions already below 90 mmHg at the index time are excluded."
    )
    add_body(
        doc,
        "The endpoint is described as blood-pressure-defined IDH because HEMOBP does not contain the symptoms and intervention data required to reconstruct a complete clinical IDH definition. Two alternative labels are retained for sensitivity analysis: a systolic decrease of at least 20 mmHg and a baseline-adjusted nadir threshold."
    )
    doc.add_heading("Analysis Unit and Prediction Boundary", level=2)
    boundary_rows = [
        ["Unit", "One linked patient-day haemodialysis session"],
        ["Prediction time", "Earliest active-dialysis record from minute 0 through minute 30"],
        ["Primary target", "Any later SBP below 90 mmHg"],
        ["Eligibility", "Index SBP at least 90 mmHg, at least two later distinct minutes, and observation through at least minute 120"],
        ["Leakage rule", "No measurement or outcome recorded after the index time is used as a predictor"],
        ["Validation grouping", "All sessions from a patient remain within one data split"],
    ]
    add_table(doc, ["Element", "Operational definition"], boundary_rows, widths=[1.55, 5.35], font_size=9.2)

    add_page_break(doc)
    doc.add_heading("2 Cohort Construction", level=1)
    add_body(
        doc,
        f"The source dataset contains {flow['patient_rows']:,} patient rows, {flow['d1_rows']:,} session-summary rows, and {flow['vip_rows']:,} time-stamped monitoring records. Patient identifier and calendar date link the session and monitoring tables. Of {flow['d1_valid_unique_patient_days']:,} valid unique patient-days, {flow['linked_patient_days']:,} linked to monitoring data. The index and follow-up rules produced {flow['final_eligible_sessions']:,} sessions from {flow['unique_patients_final']:,} patients."
    )
    add_figure(
        doc,
        ARTIFACTS / "cohort_flow.png",
        "Figure 1. Cohort construction from valid session dates to the final eligible analytic cohort.",
        width=5.85,
        alt_text="Horizontal bars show 165,849 valid patient-days, 110,790 linked patient-days, 108,781 sessions with an index record, and 106,758 final sessions.",
    )
    flow_rows = [[row.stage, f"{int(row.records_or_sessions):,}", row.action] for row in cleaning.itertuples(index=False)]
    add_table(doc, ["Stage", "Count", "Action"], flow_rows, widths=[1.75, 1.05, 4.1], font_size=7.8, vertical_margin=55)

    add_page_break(doc)
    doc.add_heading("3 Data Cleaning", level=1)
    doc.add_heading("Dates Linkage and Duplicates", level=2)
    add_body(
        doc,
        "The session table contained 128 rows without usable dates and nine duplicate patient-day keys. Invalid-date rows were excluded, and stable chronological sorting retained the first row for duplicate patient-day keys. Monitoring records were then linked many-to-one to the deduplicated session frame. No exact duplicate monitoring rows remained after linkage and physiologic screening."
    )
    doc.add_heading("Physiologic and Timing Bounds", level=2)
    add_bullets(
        doc,
        [
            "Monitoring SBP was restricted to 31 through 200 mmHg, DBP to 30 through 192 mmHg, and elapsed dialysis time to 0 through 370 minutes.",
            "The index record required active blood flow, elapsed time from 0 through 30 minutes, SBP from 60 through 200 mmHg, and DBP from 30 through 150 mmHg.",
            "Primary-outcome eligibility required index SBP of at least 90 mmHg and adequate later observation, preventing already-present hypotension and truncated follow-up from being treated as prediction cases.",
        ],
    )
    doc.add_heading("Missing Values", level=2)
    add_body(
        doc,
        "Dialysis vintage was missing in 1.37% of eligible sessions. The three prior-session features were missing in 0.78%, corresponding primarily to each patient's first eligible observed session. All other modelling fields were complete in the eligible table. Numeric values are median-imputed and categorical values are mode-imputed inside the fitted training pipeline, so validation and test information cannot influence imputation parameters."
    )
    add_figure(
        doc,
        ARTIFACTS / "missingness.png",
        "Figure 2. Percentage of missing values among eligible sessions before model-pipeline imputation.",
        width=6.3,
        alt_text="Missingness is 1.37 percent for dialysis vintage and 0.78 percent for each prior-session history feature.",
    )

    add_page_break(doc)
    doc.add_heading("4 Outlier Assessment", level=1)
    add_body(
        doc,
        "Outlier handling separates impossible measurements from unusual but potentially valid clinical observations. Hard physiologic and timing bounds remove invalid monitoring values before the index and outcome are defined. For the resulting predictors, the 1.5-IQR rule is used only as a descriptive flag. Flagged values are not automatically removed because narrow-distribution variables, such as dialysate conductivity, can produce many statistical flags that remain clinically possible."
    )
    add_body(
        doc,
        "Tree-based models use the screened values directly. Linear and PCA candidates receive imputation and standardisation within each training fold. The report retains the full outlier audit in artifacts/outlier_summary.csv so that alternative clipping or winsorisation can be evaluated as a sensitivity analysis without silently changing the primary cohort."
    )
    add_figure(
        doc,
        ARTIFACTS / "outlier_summary.png",
        "Figure 3. Descriptive IQR flags after physiologic screening. These flags identify distribution tails and are not automatic exclusion rules.",
        width=6.45,
        alt_text="Horizontal bars show the percentage outside 1.5 IQR fences for the twelve most frequently flagged predictors.",
    )

    add_page_break(doc)
    doc.add_heading("5 Outcome and Prediction Timing", level=1)
    add_body(
        doc,
        f"The primary outcome occurred in {int(df.idh_absolute.sum()):,} of {len(df):,} sessions ({df.idh_absolute.mean():.2%}). The baseline-adjusted threshold occurred in {int(df.idh_flythe.sum()):,} sessions ({df.idh_flythe.mean():.2%}), while a decrease of at least 20 mmHg occurred in {int(df.idh_drop_20.sum()):,} sessions ({df.idh_drop_20.mean():.2%}). The large difference confirms that the event rate depends strongly on the operational definition and supports retaining multiple sensitivity outcomes."
    )
    add_figure(
        doc,
        ARTIFACTS / "outcome_definition_sensitivity.png",
        "Figure 4. Outcome prevalence under the primary and two sensitivity definitions.",
        width=6.35,
        alt_text="The SBP below 90 definition occurs in 8.5 percent, the baseline-adjusted definition in 9.8 percent, and a 20 mmHg drop in 48.5 percent of sessions.",
    )
    add_body(
        doc,
        f"The median index time was {df.index_minute.median():.0f} minutes. {((df.index_minute == 0).mean()):.1%} of sessions were indexed at minute 0 and {((df.index_minute <= 5).mean()):.1%} by minute 5. The 30-minute allowance therefore covers a small tail of sessions rather than defining the typical prediction time."
    )
    add_figure(
        doc,
        ARTIFACTS / "index_time_distribution.png",
        "Figure 5. Distribution of elapsed treatment time at the selected index record.",
        width=6.45,
        alt_text="Most index records occur at minute zero, with a small tail extending to 30 minutes.",
    )

    add_page_break(doc)
    doc.add_heading("6 Applied Exploratory Data Analysis", level=1)
    doc.add_heading("Predictor Distributions", level=2)
    add_body(
        doc,
        "Outcome-stratified distributions show the clearest separation for previous-session nadir SBP and prior IDH rate. Sessions with the primary outcome also tend to have lower baseline SBP, greater fluid excess, and higher initial weight-normalised ultrafiltration. The plotted range is limited to the 0.5th through 99.5th percentiles for visual readability; the analysis table retains all values that pass physiologic screening."
    )
    add_figure(
        doc,
        ARTIFACTS / "feature_distributions.png",
        "Figure 6. Selected predictor distributions stratified by the primary outcome.",
        width=6.65,
        alt_text="Six density plots compare baseline SBP, fluid excess, initial UF rate, prior nadir SBP, prior IDH rate, and age by outcome.",
    )

    add_page_break(doc)
    doc.add_heading("Nonlinear Relationships", level=2)
    add_body(
        doc,
        "Quantile binning was used for EDA only and was not supplied to the final model. The lowest previous-session nadir decile had a 37.5% observed event rate, compared with 1.0% in the highest decile. Event prevalence rose from 4.8% to 16.3% across fluid-excess deciles and reached 15.9% in the highest initial ultrafiltration-rate decile. Baseline SBP showed a nonlinear relationship, with the greatest prevalence in the lowest decile."
    )
    add_figure(
        doc,
        ARTIFACTS / "binned_relationships.png",
        "Figure 7. Observed primary-outcome prevalence across predictor deciles with approximate 95% binomial intervals.",
        width=6.45,
        alt_text="Four plots show outcome prevalence by deciles of baseline SBP, fluid excess, initial UF rate, and previous-session nadir SBP.",
    )
    add_body(
        doc,
        "The decile plots are descriptive. They do not adjust for patient clustering, treatment decisions, or correlated predictors and must not be interpreted as causal dose-response relationships."
    )

    add_page_break(doc)
    doc.add_heading("Correlations and Subgroup Patterns", level=2)
    add_body(
        doc,
        "Spearman correlation was chosen because several predictors were skewed and relationships were not strictly linear. Previous-session nadir SBP had the largest absolute univariable correlation with the outcome (rho = -0.320), followed by prior IDH rate (rho = 0.288), fluid excess percentage (rho = 0.127), and baseline SBP (rho = -0.117). Correlated haemodynamic features were retained for model comparison, then examined through regularisation, embedded selection, and permutation analysis."
    )
    add_figure(
        doc,
        ARTIFACTS / "correlation_heatmap.png",
        "Figure 8. Spearman correlation matrix for selected predictors and the primary outcome.",
        width=5.65,
        alt_text="Heatmap shows modest correlations among selected predictors and stronger outcome relationships for prior nadir SBP and prior IDH rate.",
    )
    add_page_break(doc)
    doc.add_heading("Subgroup Patterns", level=2)
    add_body(
        doc,
        "Observed prevalence varied by recorded sex, age group, and diabetes status. These descriptive differences motivated the later fairness audit; they do not establish unfair model treatment or biological causation. Repeated sessions also give frequent attenders more influence than patients with fewer observations, so patient-grouped validation remains essential."
    )
    add_figure(
        doc,
        ARTIFACTS / "subgroup_prevalence.png",
        "Figure 9. Observed BP-defined IDH prevalence by recorded sex, age group, and diabetes status.",
        width=6.65,
        alt_text="Bar charts show higher prevalence among female sessions, ages 55 to 64, and sessions with diabetes.",
    )

    add_page_break(doc)
    doc.add_heading("7 Feature Engineering", level=1)
    add_body(
        doc,
        "Feature engineering used information available by the index time and earlier sessions only. Domain-derived features express volume burden, haemodynamic state, treatment intensity, dialysis exposure, and patient-specific history. Outcome fields, later blood pressures, post-dialysis weight, dialysis end time, identifiers, and audit variables are excluded from the model feature list."
    )
    feature_rows = [[r.engineered_feature, r.formula, r.rationale] for r in feature_engineering.itertuples(index=False)]
    add_table(doc, ["Feature", "Definition", "Rationale"], feature_rows, widths=[1.6, 2.6, 2.7], font_size=7.4, vertical_margin=55)
    add_body(
        doc,
        "History variables are calculated after chronological sorting within each patient. A one-session shift ensures that the current outcome cannot enter its own predictors. Because the train, validation, and test splits are patient-disjoint, no patient's history crosses a split boundary."
    )

    add_page_break(doc)
    doc.add_heading("8 Preprocessing Pipeline", level=1)
    preprocessing_rows = [
        ["Numeric missing values", "Median imputation", "Fitted inside each training fold"],
        ["Categorical missing values", "Most-frequent imputation", "Fitted inside each training fold"],
        ["Recorded sex and diabetes", "One-hot encoding", "Unknown categories ignored safely"],
        ["Numeric linear-model inputs", "Standard scaling", "Required for regularised logistic regression and PCA"],
        ["Tree-model inputs", "No scaling", "Preserves natural split thresholds"],
        ["Prior session count", "log(1 + count)", "Reduces strong right skew"],
        ["EDA relationship plots", "Quantile bins", "Diagnostic only; not model inputs"],
        ["Outliers", "Physiologic bounds plus retention of plausible extremes", "Avoids arbitrary deletion of rare clinical observations"],
    ]
    add_table(doc, ["Input", "Transformation", "Justification"], preprocessing_rows, widths=[1.7, 2.15, 3.05], font_size=8.4)
    doc.add_heading("Leakage Controls", level=2)
    add_bullets(
        doc,
        [
            "Only later records define the outcome; later measurements are never candidate predictors.",
            "Imputation, scaling, one-hot encoding, L1 selection, and PCA are fitted within scikit-learn pipelines rather than before cross-validation.",
            "Cross-validation uses patient groups, preventing sessions from the same patient from appearing in both fitting and validation folds.",
            "Feature importance on the test set is used only for final explanation, not for tuning or feature selection.",
        ],
    )
    doc.add_heading("Clustering Tendency", level=2)
    add_body(
        doc,
        "Clustering tendency was not evaluated because DIAL-ALERT is a supervised classification task with a prespecified outcome. Adding K-means, DBSCAN, t-SNE, or UMAP would not answer the stated prediction question. PCA was retained as the required dimensionality-reduction comparison."
    )

    add_page_break(doc)
    doc.add_heading("9 Embedded Feature Selection", level=1)
    add_body(
        doc,
        f"An embedded selection pipeline used L1-penalised logistic regression through SelectFromModel. The selector retained {int(selected.selected.sum())} of {len(selected)} transformed features using the median importance threshold. Selection occurred within grouped cross-validation, preventing the holdout test set from determining the feature set."
    )
    readable_selected = [name.replace("_", " ") for name in selected_names]
    selected_rows = [readable_selected[i:i + 2] for i in range(0, len(readable_selected), 2)]
    add_table(doc, ["Selected feature", "Selected feature"], selected_rows, widths=[3.45, 3.45], font_size=9.0)
    add_body(
        doc,
        "The L1 pipeline provides explicit evidence of feature selection but is treated as a candidate model rather than an automatic filter for every algorithm. Gradient boosting can use nonlinearities and correlated signals differently, so its inputs were not removed based on a linear selector. Test-set permutation importance was also excluded from selection decisions."
    )

    add_page_break(doc)
    doc.add_heading("10 Principal Component Analysis", level=1)
    add_body(
        doc,
        f"PCA was applied only to numeric predictors after median imputation and standardisation. The selected PCA logistic-regression candidate used an 85% variance threshold and retained {len(pca)} components, which explained {pca.cumulative_explained_variance.iloc[-1]:.1%} of numeric-feature variance. Categorical variables remained one-hot encoded outside the PCA block."
    )
    add_figure(
        doc,
        ARTIFACTS / "pca_variance.png",
        "Figure 10. Cumulative explained variance for the PCA candidate selected during grouped cross-validation.",
        width=6.3,
        alt_text="Cumulative variance rises to 87.6 percent after eleven principal components and crosses the 85 percent threshold.",
    )
    add_body(
        doc,
        "PCA logistic regression achieved grouped cross-validation average precision of 0.410, compared with 0.446 for histogram gradient boosting. PCA was therefore documented as a dimensionality-reduction benchmark rather than selected for the final predictor. Its lower performance and reduced feature-level interpretability outweighed its compression benefit for this dataset."
    )

    add_page_break(doc)
    doc.add_heading("11 Global Feature Importance", level=1)
    add_body(
        doc,
        "Permutation importance measures the reduction in test-set average precision after a feature is shuffled. Previous-session nadir SBP produced the largest decrease (0.128), followed by prior IDH rate (0.081), fluid excess percentage (0.015), baseline SBP (0.009), and baseline mean arterial pressure (0.004). Near-zero or negative values can occur when predictors are redundant or when finite-sample variation exceeds the feature's unique contribution."
    )
    add_figure(
        doc,
        ARTIFACTS / "permutation_importance.png",
        "Figure 11. Test-set permutation importance with variability across repeated shuffles.",
        width=6.0,
        alt_text="Prior nadir SBP and prior IDH rate dominate the decrease in average precision after permutation.",
    )
    top_rows = [
        [row.feature.replace("_", " "), f"{row.importance_mean:.3f}", f"{row.importance_sd:.3f}"]
        for row in importance.head(7).itertuples(index=False)
    ]
    add_table(doc, ["Feature", "Mean decrease", "Standard deviation"], top_rows, widths=[3.8, 1.45, 1.65], font_size=8.8)

    add_page_break(doc)
    doc.add_heading("12 SHAP and Conditional Effects", level=1)
    add_body(
        doc,
        "SHAP explanations were generated for the uncalibrated base model because probability calibration does not change the underlying predictor relationships. Only aggregate mean absolute contributions are published. The summary confirms that previous-session nadir pressure, prior IDH rate, and fluid excess dominate model output. SHAP values describe the fitted model and do not prove that modifying a feature would change a patient's outcome."
    )
    add_figure(
        doc,
        ARTIFACTS / "step5_shap_summary.png",
        "Figure 12. Aggregate approximate SHAP importance for the locked model.",
        width=6.05,
        alt_text="Aggregate SHAP bar chart shows prior nadir SBP, prior IDH rate, and fluid excess percent as dominant model contributors.",
    )
    add_body(
        doc,
        "Synthetic-reference-profile dependence curves provide a second check on direction and heterogeneity without publishing patient-level trajectories. Average predicted dependence decreases as previous nadir SBP rises and increases as prior IDH rate or fluid excess rises. Variation across the predefined synthetic profiles illustrates model interactions rather than observed patient histories."
    )
    add_figure(
        doc,
        ARTIFACTS / "step5_pdp_ice.png",
        "Figure 13. Dependence and conditional-effect curves generated from synthetic reference profiles.",
        width=6.65,
        alt_text="Three panels show decreasing model dependence for prior nadir SBP and increasing dependence for prior IDH rate and fluid excess percent.",
    )

    doc.add_heading("13 Limitations and Mitigations", level=1)
    limitations = [
        ["BP-defined endpoint", "Symptoms and interventions are unavailable, so the study cannot reconstruct the complete clinical syndrome.", "Use BP-defined terminology and report alternative threshold definitions."],
        ["Informative measurement", "Unstable patients may receive more frequent BP measurements and have more opportunities for event detection.", "Report this ascertainment risk and consider fixed-time-window sensitivity analyses."],
        ["Single-centre historical data", "Practice patterns, equipment, and patient mix may not represent other countries or current care.", "Require external and prospective validation before clinical use."],
        ["Linked subset", "Only patient-days with linkable detailed monitoring enter the final cohort.", "Show the cohort flow and compare excluded sessions when source fields permit."],
        ["Repeated sessions", "The effective sample size is smaller than the number of sessions.", "Use patient-grouped splits and patient-clustered uncertainty estimates."],
        ["History dependence", "New patients lack prior-session predictors.", "Use pipeline imputation and report separate cold-start performance in later analyses."],
        ["Residual extreme values", "Statistical outliers may reflect error or true clinical extremes.", "Retain plausible values, publish the audit, and test clipping as a sensitivity analysis."],
        ["Variable index time", "A small minority of predictions occur after minute 5.", "Report timing and evaluate fixed prediction windows."],
    ]
    add_table(doc, ["Issue", "Possible effect", "Mitigation"], limitations, widths=[1.35, 2.8, 2.75], font_size=7.8)
    add_body(
        doc,
        "These limitations do not invalidate the capstone. They define the scope of the evidence: retrospective model development and internal validation of a BP-defined operational target. Clinical effectiveness, treatment benefit, and cost savings remain prospective research questions."
    )

    add_page_break(doc)
    doc.add_heading("14 Reproducibility", level=1)
    add_body(
        doc,
        "The analysis uses configuration-driven feature lists and random state 42. All patient grouping, transformations, feature selection, PCA, and model fitting are implemented in reproducible Python scripts. Raw HEMOBP files should be obtained from the cited Figshare record and placed in data/raw."
    )
    doc.add_heading("Execution Commands", level=2)
    add_code(
        doc,
        [
            "python src/build_session_dataset.py --raw-dir data/raw --output-dir data/processed",
            "python src/train_evaluate.py --data data/processed/hemobp_session_level.csv.gz",
            "python src/generate_eda.py",
            "python src/create_eda_report.py",
        ],
    )
    doc.add_heading("Primary Reproducibility Files", level=2)
    files_rows = [
        ["src/build_session_dataset.py", "Cleaning, linkage, index-time definition, outcomes, and domain features"],
        ["src/generate_eda.py", "EDA figures, cleaning audit, outlier audit, and dictionaries"],
        ["src/train_evaluate.py", "Grouped splitting, preprocessing, feature selection, PCA, explainability, and models"],
        ["configs/model_config.json", "Outcome, features, alert capacity, and random seed"],
        ["docs/data_dictionary.csv", "Complete 40-variable processed-table dictionary"],
        ["docs/feature_engineering_dictionary.csv", "Engineered-feature formula and rationale"],
        ["artifacts/l1_selected_features.csv", "Embedded-selection output"],
        ["artifacts/pca_explained_variance.csv", "PCA variance evidence"],
        ["artifacts/permutation_importance.csv", "Global importance values and uncertainty"],
        ["artifacts/outlier_summary.csv", "Descriptive outlier flags and disposition"],
    ]
    add_table(doc, ["File", "Purpose"], files_rows, widths=[2.55, 4.35], font_size=8.2)
    doc.add_heading("Quality Checks", level=2)
    add_bullets(
        doc,
        [
            "Processed row count, event count, patient count, and date range are recorded in cohort_flow.json.",
            "Pipeline transformations are fitted only on training folds and tolerate unseen categorical levels.",
            "Generated tables preserve numeric values used in every figure, allowing visual claims to be checked directly.",
            "The final report is rendered to page images and inspected before submission.",
        ],
    )

    add_page_break(doc)
    doc.add_heading("15 Conclusions", level=1)
    add_body(
        doc,
        "The Step 3 analysis produced a documented, reproducible, and leakage-controlled session table suitable for supervised model development. Missing data were uncommon and are handled within training pipelines. Invalid readings were removed through explicit bounds, while unusual but plausible values remained visible through a published outlier audit. Domain engineering added clinically interpretable volume, haemodynamic, treatment-intensity, and longitudinal-history predictors."
    )
    add_body(
        doc,
        "EDA consistently identified prior haemodynamic instability as the strongest predictive signal, with additional information from fluid excess and baseline blood pressure. Embedded L1 selection and PCA satisfy complementary reduction objectives, but neither was allowed to use the test set for selection. Permutation importance, aggregate SHAP, and synthetic-profile dependence curves provide consistent post hoc explanations while preserving the distinction between prediction and causation."
    )
    add_body(
        doc,
        "The resulting evidence supports progression to formal model comparison, calibration, subgroup auditing, and operational threshold evaluation. It does not establish clinical benefit or readiness for deployment."
    )

    doc.add_heading("References", level=1)
    refs = [
        "1. Lin CJ, Chen YY, Pan CF, Wu VC, Wu CJ. Dataset supporting blood pressure prediction for the management of chronic hemodialysis. Scientific Data. 2019;6:313. https://doi.org/10.1038/s41597-019-0319-8",
        "2. Lin CJ and colleagues. Hemrec VIP csv. Figshare dataset version 3. https://doi.org/10.6084/m9.figshare.6260654.v3",
        "3. Lundberg SM, Lee SI. A unified approach to interpreting model predictions. Advances in Neural Information Processing Systems. 2017;30.",
        "4. Jolliffe IT, Cadima J. Principal component analysis a review and recent developments. Philosophical Transactions of the Royal Society A. 2016;374:20150202. https://doi.org/10.1098/rsta.2015.0202",
        "5. Pedregosa F and colleagues. Scikit-learn machine learning in Python. Journal of Machine Learning Research. 2011;12:2825-2830.",
        "6. Flythe JE, Xue H, Lynch KE, Curhan GC, Brunelli SM. Association of mortality risk with various definitions of intradialytic hypotension. Journal of the American Society of Nephrology. 2015;26:724-734. https://doi.org/10.1681/ASN.2014020222",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(-0.2)
        p.paragraph_format.left_indent = Inches(0.2)
        r = p.add_run(ref)
        set_font(r, size=8.8)

    doc.core_properties.title = "DIAL ALERT EDA and Feature Engineering Report"
    doc.core_properties.subject = "Capstone Project Step 3"
    doc.core_properties.author = "Franklin Guillano"
    doc.core_properties.keywords = "DIAL-ALERT, HEMOBP, EDA, feature engineering, haemodialysis, machine learning"
    doc.core_properties.comments = "DIAL-ALERT academic capstone"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(make_report())
