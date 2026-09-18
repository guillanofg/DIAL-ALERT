#!/usr/bin/env python3
"""Create the DIAL ALERT Step 2 data collection and understanding report."""

from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS = ROOT / "artifacts"
REPORTS = ROOT / "reports"
OUT = REPORTS / "Franklin_Guillano_DIAL_ALERT_Data_Collection_and_Understanding_Report.docx"

NAVY = "17324D"
BLUE = "1D6F8A"
TEAL = "159A9C"
PALE_BLUE = "EAF3F7"
PALE_TEAL = "E7F5F3"
PALE_GOLD = "FFF3D6"
LIGHT_GRAY = "F2F4F6"
MID_GRAY = "D5DDE3"
TEXT = RGBColor(32, 43, 54)
MUTED = RGBColor(89, 104, 117)
WHITE = RGBColor(255, 255, 255)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=80, bottom=70, end=80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=MID_GRAY, size="4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def style_run(run, size=None, bold=None, color=TEXT, italic=None, font="Arial") -> None:
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = color


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    title.font.size = Pt(29)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(NAVY)
    title.paragraph_format.space_after = Pt(10)
    # Word's built-in Title style can carry a bottom border; remove it so the
    # title is typographic rather than decorated with an automatic rule.
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size, color, before, after in (
        ("Heading 1", 18, NAVY, 12, 7),
        ("Heading 2", 13, BLUE, 10, 5),
        ("Heading 3", 10.5, NAVY, 7, 3),
    ):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    caption.font.name = "Arial"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    caption.font.size = Pt(8)
    caption.font.italic = True
    caption.font.color.rgb = MUTED
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(7)


def configure_section(section, landscape=False, first=False) -> None:
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11)
        section.page_height = Inches(8.5)
        section.left_margin = Inches(0.55)
        section.right_margin = Inches(0.55)
        section.top_margin = Inches(0.58)
        section.bottom_margin = Inches(0.65)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.left_margin = Inches(0.72)
        section.right_margin = Inches(0.72)
        section.top_margin = Inches(0.68)
        section.bottom_margin = Inches(0.78)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.3)
    section.different_first_page_header_footer = bool(first)


def configure_headers_and_footers(doc: Document) -> None:
    first = doc.sections[0]
    header = first.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    p.text = "DIAL ALERT CAPSTONE STEP 2"
    style_run(p.runs[0], size=7.5, bold=True, color=RGBColor.from_string(BLUE))
    footer = first.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    fp.text = ""
    r = fp.add_run("Franklin Guillano   Data Collection and Understanding   ")
    style_run(r, size=7.3, color=MUTED)
    add_page_number(fp)
    for run in fp.runs[1:]:
        style_run(run, size=7.3, color=MUTED)
    for section in doc.sections[1:]:
        section.header.is_linked_to_previous = True
        section.footer.is_linked_to_previous = True


def add_body(doc, text, bold_prefix=None, italic=False, space_after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        style_run(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        style_run(r2, italic=italic)
    else:
        r = p.add_run(text)
        style_run(r, italic=italic)
    return p


def add_bullets(doc, items, level=0, font_size=9.2):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.left_indent = Inches(0.22 + 0.18 * level)
        p.paragraph_format.first_line_indent = Inches(-0.12)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        style_run(r, size=font_size)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.24)
        p.paragraph_format.first_line_indent = Inches(-0.14)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        style_run(r, size=9.2)


def add_band(doc, text, fill=PALE_BLUE, color=NAVY):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.95)
    cell = table.cell(0, 0)
    shade(cell, fill)
    set_cell_margins(cell, top=105, bottom=105, start=120, end=120)
    remove_table_borders(table)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    style_run(r, size=9.2, bold=True, color=RGBColor.from_string(color))
    return table


def add_table(doc, headers, rows, widths, font_size=8.2, header_fill=NAVY, alternate=True):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    repeat_header(header)
    prevent_split(header)
    for i, (label, width) in enumerate(zip(headers, widths)):
        header.cells[i].width = Inches(width)
        shade(header.cells[i], header_fill)
        set_cell_margins(header.cells[i], top=80, bottom=80, start=70, end=70)
        header.cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = header.cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(label))
        style_run(r, size=font_size, bold=True, color=WHITE)
    for ridx, row_data in enumerate(rows):
        row = table.add_row()
        prevent_split(row)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for i, (value, width) in enumerate(zip(row_data, widths)):
            cell = row.cells[i]
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell, top=65, bottom=65, start=65, end=65)
            if alternate and ridx % 2 == 1:
                shade(cell, "F7FAFC")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(str(value))
            style_run(r, size=font_size, color=TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_kpi_strip(doc, kpis):
    table = doc.add_table(rows=1, cols=len(kpis))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    remove_table_borders(table)
    width = 6.85 / len(kpis)
    for i, (value, label) in enumerate(kpis):
        cell = table.cell(0, i)
        cell.width = Inches(width)
        shade(cell, PALE_TEAL if i % 2 == 0 else PALE_BLUE)
        set_cell_margins(cell, top=115, bottom=115, start=70, end=70)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(value)
        style_run(r, size=17, bold=True, color=RGBColor.from_string(NAVY))
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(label)
        style_run(r2, size=7.8, bold=True, color=MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure(doc, image_path: Path, caption: str, width: float, page_break=False):
    if page_break:
        doc.add_page_break()
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    remove_table_borders(table)
    row = table.rows[0]
    prevent_split(row)
    cell = row.cells[0]
    set_cell_margins(cell, top=0, bottom=0, start=0, end=0)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    p.add_run().add_picture(str(image_path), width=Inches(width))
    cp = cell.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.style = doc.styles["Caption"]
    cr = cp.add_run(caption)
    style_run(cr, size=8, color=MUTED, italic=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def pct(n, d):
    return f"{100*n/d:.3f}%"


def missing_text(row):
    n = int(row["missing_n"])
    p = float(row["missing_pct"])
    return f"{n:,} ({p:.3f}%)"


def cover_page(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(34)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("CAPSTONE PROJECT STEP 2")
    style_run(r, size=10, bold=True, color=RGBColor.from_string(TEAL))

    p = doc.add_paragraph(style="Title")
    p.add_run("DIAL ALERT Data Collection and Understanding")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(26)
    r = p.add_run("Dataset overview and complete data dictionary")
    style_run(r, size=16, color=RGBColor.from_string(BLUE))

    add_kpi_strip(
        doc,
        [
            ("1,072", "PUBLIC PATIENT RECORDS"),
            ("165,986", "SESSION ROWS"),
            ("4.37M", "MONITOR RECORDS"),
            ("106,758", "ELIGIBLE SESSIONS"),
        ],
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(20)
    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta.autofit = False
    remove_table_borders(meta)
    labels = ["Prepared by", "Project", "Source dataset", "Report date"]
    values = [
        "Franklin Guillano",
        "DIAL ALERT Early Warning for Intradialytic Hypotension",
        "HEMOBP Version 3",
        "September 2026",
    ]
    for row, label, value in zip(meta.rows, labels, values):
        row.cells[0].width = Inches(1.25)
        row.cells[1].width = Inches(5.35)
        for cell in row.cells:
            set_cell_margins(cell, top=35, bottom=35, start=0, end=70)
        p1 = row.cells[0].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        rr = p1.add_run(label.upper())
        style_run(rr, size=7.5, bold=True, color=MUTED)
        p2 = row.cells[1].paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        rv = p2.add_run(value)
        style_run(rv, size=9.5, bold=(label == "Prepared by"))

    doc.add_paragraph().paragraph_format.space_after = Pt(22)
    add_band(doc, "ACADEMIC DECISION SUPPORT STUDY   NOT FOR CLINICAL CARE", fill=PALE_GOLD, color=NAVY)
    add_body(
        doc,
        "This report documents dataset provenance, structure, completeness, validity checks, cohort construction, and variable definitions before model development. It does not establish clinical safety or readiness for patient care.",
        italic=True,
        space_after=0,
    )
    doc.add_page_break()


def build_report():
    REPORTS.mkdir(parents=True, exist_ok=True)
    audit = json.loads((ARTIFACTS / "step2_dataset_audit.json").read_text())
    dictionary = json.loads((ARTIFACTS / "step2_dictionary_rows.json").read_text())
    raw_rows = [r for r in dictionary if r["dataset_layer"] == "Raw source"]
    analytic_rows = [r for r in dictionary if r["dataset_layer"] == "Analytic session table"]

    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0], first=True)
    cover_page(doc)

    doc.add_heading("Executive Summary", level=1)
    add_body(
        doc,
        "DIAL ALERT uses HEMOBP Version 3, a public, de-identified hemodialysis dataset released on Figshare under CC BY 4.0 and described in Scientific Data. The source combines patient characteristics, session-level treatment records, and time-stamped blood-pressure and dialysis-machine measurements. This longitudinal structure is well aligned with the proposed classification task: estimate, at the start of a dialysis session, the probability of a later blood-pressure-defined intradialytic hypotension event. [1,2]",
    )
    add_body(
        doc,
        "The downloaded public files contain 1,072 patient records, 165,986 session rows, and 4,366,298 monitor records. After deterministic patient-date linkage and preregistered eligibility checks, the analysis-ready table contains 106,758 sessions from 830 patients. The primary label occurs in 9,067 sessions, a prevalence of 8.493%. All counts in this report come from the actual files, not from values copied from the publication.",
    )
    add_body(
        doc,
        "The dataset is strong enough for a rigorous academic capstone, but it is not a deployment-grade clinical dataset. It is retrospective and single-center; the public files do not contain symptoms, interventions, medications, laboratory results, or broad comorbidity profiles. Accordingly, the primary outcome is explicitly described as BP-defined IDH rather than a complete symptomatic clinical diagnosis.",
    )
    add_kpi_strip(
        doc,
        [
            ("830", "ANALYTIC PATIENTS"),
            ("40", "ANALYTIC VARIABLES"),
            ("8.493%", "PRIMARY EVENT RATE"),
            ("CC BY", "DATA LICENSE"),
        ],
    )

    doc.add_heading("Rubric Evidence", level=2)
    evidence = [
        ["High-quality dataset and source cited", "Peer-reviewed data descriptor, DOI-stable Figshare release, license documented, and file hashes verified."],
        ["Comprehensive dataset overview", "Grain, dimensions, date coverage, linkage keys, cohort flow, missingness, duplicates, ranges, and outlier policy are reported."],
        ["Complete data dictionary", "All 23 raw fields and all 40 analysis-table fields are defined with type, unit, values, missingness, meaning, and analytic role."],
        ["Reproducibility", "Raw files are immutable; transformation rules are code-based; checksums, paths, and commands are documented."],
    ]
    add_table(doc, ["Rubric expectation", "Evidence in this deliverable"], evidence, [2.0, 4.85], font_size=8.4)

    doc.add_heading("Dataset Selection and Justification", level=1)
    add_body(
        doc,
        "HEMOBP was selected because its three linked tables support the exact temporal structure required by DIAL ALERT. Patient characteristics and predialysis measurements are available before or at the index time, while subsequent blood-pressure observations can define the outcome. This separation permits an auditable prediction-time boundary and reduces the risk of post-index data leakage.",
    )
    add_bullets(
        doc,
        [
            "Industry relevance: hemodialysis is a repeated, safety-critical treatment process in which hypotension risk must be assessed session by session.",
            "Modelling relevance: the data support supervised binary classification, longitudinal history features, class-imbalance analysis, explainability, and patient-grouped validation.",
            "Scale: more than four million monitor records provide enough repeated observations for robust preprocessing and model comparison.",
            "Transparency: the peer-reviewed descriptor and versioned public archive make the acquisition route, license, and data-generating context traceable.",
        ],
    )

    doc.add_page_break()
    doc.add_heading("Source and Provenance", level=1)
    source_rows = []
    for item in audit["tables"]:
        source_rows.append(
            [
                item["file"],
                item["grain"],
                f'{item["rows"]:,} x {item["columns"]}',
                item["date_range"],
                f'{item["size_bytes"] / (1024*1024):.2f} MB',
                item["observed_md5"],
            ]
        )
    add_table(
        doc,
        ["File", "Grain", "Rows x columns", "Coverage", "Size", "Verified MD5"],
        source_rows,
        [0.72, 1.35, 0.9, 1.35, 0.72, 1.8],
        font_size=7.3,
    )
    add_body(
        doc,
        "Integrity result. Each local MD5 checksum exactly matches the checksum published in the Figshare Version 3 metadata. This verifies that the audited files are byte-for-byte copies of the cited public release.",
        bold_prefix="Integrity result.",
    )
    add_body(
        doc,
        "Source-count discrepancy. The Scientific Data article reports 1,075 outpatients, whereas the released Idp file contains 1,072 rows and 1,072 unique patient identifiers, a difference of three. No undocumented records were reconstructed. All analyses use the public file contents, and the discrepancy remains an explicit provenance limitation. [1,2]",
        bold_prefix="Source-count discrepancy.",
    )
    doc.add_heading("Collection Context and Governance", level=2)
    context_rows = [
        ["Care setting", "Retrospective outpatient chronic hemodialysis care at a tertiary center in Taiwan"],
        ["Observation window", "June 2013 to July 2018 across the released tables"],
        ["Ethics record", "Institutional review board identifier 16MMHIS044, as reported in the data descriptor"],
        ["Privacy", "De-identified patient identifiers; dates and times retained for longitudinal linkage"],
        ["Reuse conditions", "Figshare Version 3 distributed under Creative Commons Attribution 4.0"],
    ]
    add_table(doc, ["Element", "Documented context"], context_rows, [1.55, 5.15], font_size=8.2)

    add_figure(
        doc,
        ARTIFACTS / "data_architecture.png",
        "Figure 1. Three source tables are linked by de-identified patient ID and calendar date, then collapsed to one prediction row per eligible dialysis session.",
        width=6.72,
        page_break=True,
    )
    doc.add_heading("Data Architecture and Linkage", level=1)
    add_body(
        doc,
        "The source has three different grains. Idp contains one row per patient. Hemrec D1 contains one row per recorded dialysis session. Hemrec VIP contains repeated, time-stamped monitor states within a session. The analytic unit is one eligible patient-session, not one monitor reading and not one patient.",
    )
    add_numbered(
        doc,
        [
            "Normalize D1 and VIP dates and link records on de-identified patient ID plus calendar date.",
            "Use the earliest valid active-dialysis VIP record within the first 30 elapsed minutes as the index record. Active treatment requires blood flow greater than zero.",
            "Require index SBP from 90 to 200 mmHg, index DBP from 30 to 150 mmHg, at least two later distinct elapsed-minute values, and an observation at or beyond dialysis minute 120.",
            "Derive the primary outcome only from records after the index observation. Retain outcome and post-treatment fields for audit, but exclude them from predictors.",
            "Derive prior-session features with chronological shifts and cumulative calculations so the current or future session cannot enter its own history.",
        ],
    )
    add_body(
        doc,
        "Linkage assumption. Calendar date is the shared session key available in the public tables. Duplicate patient-date D1 rows are resolved deterministically after stable ordering. If a patient had two distinct treatments on one calendar date, this key could conflate them; that risk cannot be fully tested because no public session ID is provided.",
        bold_prefix="Linkage assumption.",
    )
    add_body(
        doc,
        "Sampling pattern. VIP observations are not perfectly regular. The paper notes routine blood-pressure collection at approximately 30-minute intervals with additional measurements according to clinical need. Measurement density may therefore contain information about clinician concern and should not be treated as random. [1]",
        bold_prefix="Sampling pattern.",
    )

    doc.add_heading("Feature Type Summary", level=2)
    raw_counts = Counter(r["data_type"] for r in raw_rows)
    analytic_counts = Counter(r["data_type"] for r in analytic_rows)
    types = sorted(set(raw_counts) | set(analytic_counts))
    type_rows = [[t.title(), raw_counts.get(t, 0), analytic_counts.get(t, 0)] for t in types]
    add_table(doc, ["Data type", "Raw fields", "Analytic fields"], type_rows, [3.2, 1.75, 1.75], font_size=8.5)

    doc.add_heading("Cohort Construction", level=1)
    add_figure(
        doc,
        ARTIFACTS / "cohort_flow.png",
        "Figure 2. Deterministic cohort flow from the three public source files to 106,758 eligible session-level observations.",
        width=6.35,
    )
    cf = audit["cohort_flow"]
    cohort_rows = [
        ["Patient records in Idp", f'{cf["patient_rows"]:,}', "Patient roster available for linkage"],
        ["D1 raw session rows", f'{cf["d1_rows"]:,}', "Includes 128 rows without a usable session date"],
        ["Valid unique patient-days in D1", f'{cf["d1_valid_unique_patient_days"]:,}', "After invalid-date removal and patient-date deduplication"],
        ["VIP rows linked to D1", f'{cf["vip_rows_linked_to_d1"]:,}', "Monitor rows with a matching patient-date session"],
        ["Linked patient-days", f'{cf["linked_patient_days"]:,}', "At least one linked VIP record"],
        ["Patient-days with valid index", f'{cf["patient_days_with_index_record"]:,}', "Earliest active record within 30 minutes and valid index BP"],
        ["Final eligible sessions", f'{cf["final_eligible_sessions"]:,}', "Baseline SBP at least 90 and adequate post-index follow-up"],
    ]
    add_table(doc, ["Stage", "Count", "Rule or interpretation"], cohort_rows, [2.65, 1.0, 3.05], font_size=8.1)
    add_body(
        doc,
        "Selection implication. The final cohort includes 830 of the 1,072 public patients. Excluding baseline hypotension, short or sparsely measured sessions, and unlinked patient-days improves label validity but may make the analytic cohort systematically healthier, more observable, or operationally different from excluded sessions. External validation is therefore required before clinical use.",
        bold_prefix="Selection implication.",
    )

    doc.add_heading("Data Quality and Missingness", level=1)
    add_figure(
        doc,
        ARTIFACTS / "missingness.png",
        "Figure 3. Missingness is concentrated in 128 D1 rows and in history or dialysis-vintage variables in the final analytic table.",
        width=6.45,
    )
    missing_rows = [
        ["Idp", "All 5 variables", "0", "0.000%", "No source missingness detected"],
        ["Hemrec D1", "All fields except pid", "128 each", "0.077%", "The same 128 rows lack all session details"],
        ["Hemrec VIP", "All 10 variables", "0", "0.000%", "No encoded nulls detected"],
        ["Analytic", "dialysis_vintage_years", "1,465", "1.372%", "Unavailable or invalid chronology"],
        ["Analytic", "prior_session_idh", "830", "0.777%", "First eligible session per patient"],
        ["Analytic", "prior_nadir_sbp", "830", "0.777%", "First eligible session per patient"],
        ["Analytic", "prior_idh_rate", "830", "0.777%", "First eligible session per patient"],
        ["Analytic", "Remaining 36 variables", "0", "0.000%", "Complete after eligibility filters"],
    ]
    add_table(doc, ["Table", "Variable group", "Missing n", "Missing percent", "Interpretation"], missing_rows, [0.85, 1.55, 0.75, 0.9, 2.65], font_size=7.8)
    add_bullets(
        doc,
        [
            "Null handling is rule-based. D1 rows without a session date cannot be linked and are excluded before cohort construction.",
            "History-feature missingness is structural, not evidence of poor source quality: the first eligible session has no earlier eligible session.",
            "Dialysis-vintage missingness remains explicit and will be handled inside the modelling pipeline, fitted on training data only.",
            "No complete-case deletion is applied merely because a predictor is missing; the modelling stage will compare transparent imputation and missing-indicator strategies.",
        ],
    )

    doc.add_heading("Duplicates and Referential Integrity", level=2)
    duplicate_rows = [
        ["Idp", "0 exact duplicates; 0 duplicate patient IDs", "Patient key is unique"],
        ["Hemrec D1", "5 exact duplicates; 9 duplicate patient-date rows", "Stable order retained and first row kept for one session per patient-date"],
        ["Hemrec VIP", "0 exact duplicates after linkage and validity screening", "No linked analytic monitor row removed as an exact duplicate"],
        ["Analytic table", "One row per eligible patient-date", "Patient-date is unique by construction"],
    ]
    add_table(doc, ["Table", "Finding", "Disposition"], duplicate_rows, [1.05, 2.55, 3.1], font_size=8.1)

    doc.add_heading("Outliers and Plausibility", level=1)
    add_body(
        doc,
        "Outlier detection separates validity from rarity. Clearly unusable values are screened through predefined physiological and operational bounds before index selection. The interquartile-range rule is then used only to describe unusual but potentially valid observations; it is not used as an automatic deletion rule. This protects clinically important extremes and prevents data-dependent trimming from silently changing the target population.",
    )
    add_figure(
        doc,
        ARTIFACTS / "outlier_summary.png",
        "Figure 4. IQR flags among candidate predictors. High flag rates can reflect narrow central distributions rather than invalid records.",
        width=6.25,
    )
    outlier_rows = []
    for item in audit["outliers"][:10]:
        outlier_rows.append(
            [
                item["feature"],
                f'{item["minimum"]:.2f}',
                f'{item["median"]:.2f}',
                f'{item["maximum"]:.2f}',
                f'{item["iqr_flagged_n"]:,}',
                f'{item["iqr_flagged_pct"]:.2f}%',
            ]
        )
    add_table(doc, ["Feature", "Min", "Median", "Max", "IQR flags", "Flag rate"], outlier_rows, [2.35, 0.75, 0.8, 0.8, 0.95, 0.9], font_size=7.8)
    add_body(
        doc,
        "Priority quality checks. Conductivity has the largest IQR flag rate, 12.94%, because most readings are tightly concentrated near 14 mS/cm. Derived fluid-excess values range from -23.8 to 37.7 kg, and normalized ultrafiltration reaches 92.88 mL/kg/hour. These values are retained for transparency but will receive sensitivity checks, robust scaling or winsorization fitted only on training data, and individual-record review where feasible.",
        bold_prefix="Priority quality checks.",
    )

    doc.add_heading("Outcome Definition and Dataset Fitness", level=1)
    add_figure(
        doc,
        ARTIFACTS / "outcome_definition_sensitivity.png",
        "Figure 5. Outcome prevalence varies substantially by blood-pressure definition, reinforcing the need to preregister a primary label and report sensitivity analyses.",
        width=6.4,
    )
    outcome_rows = [
        ["Primary", "idh_absolute", "Any later SBP below 90 mmHg", f'{cf["primary_idh_events"]:,}', f'{cf["primary_idh_prevalence"]:.3%}'],
        ["Sensitivity", "idh_flythe", "Later nadir below 90 if baseline below 160; otherwise below 100", f'{cf["flythe_events"]:,}', f'{cf["flythe_events"]/cf["final_eligible_sessions"]:.3%}'],
        ["Sensitivity", "idh_drop_20", "Any later SBP at least 20 mmHg below baseline", f'{cf["drop_20_events"]:,}', f'{cf["drop_20_events"]/cf["final_eligible_sessions"]:.3%}'],
    ]
    add_table(doc, ["Status", "Variable", "Operational definition", "Events", "Prevalence"], outcome_rows, [0.85, 1.05, 3.25, 0.8, 0.85], font_size=7.7)
    add_body(
        doc,
        "Fitness conclusion. HEMOBP is fit for an academic study of early, BP-defined risk prediction. It is not sufficient to claim prediction of the full clinical syndrome of IDH because symptoms and interventions are absent. It also does not support causal claims about treatment settings: dialysate temperature, conductivity, ultrafiltration, and blood flow were assigned in routine care rather than randomized.",
        bold_prefix="Fitness conclusion.",
    )

    doc.add_heading("Known Limitations and Controls", level=2)
    limitations = [
        ["Single center and country", "Transportability to other dialysis units, devices, and populations is unknown.", "Use patient-grouped internal validation and require independent external validation."],
        ["Repeated sessions", "A random row split would place the same patient in training and testing.", "Use patient-grouped splits and report patient-level bootstrap uncertainty."],
        ["Outcome incompleteness", "No symptoms or rescue interventions are available.", "Name the label BP-defined IDH and present alternate BP definitions."],
        ["Informative sampling", "Extra BP readings may occur when clinicians are concerned.", "Exclude post-index count features from prediction and discuss measurement-process bias."],
        ["Limited clinical covariates", "Most laboratory, medication, and comorbidity data are absent.", "Constrain claims and document omitted-variable risk."],
        ["Sensitive attributes", "Public data include recorded sex and age but not race or socioeconomic status.", "Audit available subgroups and state that unobserved-group fairness cannot be assessed."],
    ]
    add_table(doc, ["Limitation", "Why it matters", "Planned control"], limitations, [1.35, 2.55, 2.8], font_size=7.7)

    doc.add_heading("Reproducibility Record", level=1)
    add_body(
        doc,
        "The source CSV files remain immutable in data/raw. The session-level table is generated by code, and every exclusion and feature definition is encoded rather than performed manually. The workbook supplied with this report contains the full audit fields and can be filtered by source table, modelling role, or prediction-time availability.",
    )
    commands = [
        "python src/build_session_dataset.py",
        "python src/generate_step2_assets.py",
        "node src/create_step2_dictionary_workbook.mjs",
        "python src/create_step2_report.py",
    ]
    add_table(doc, ["Purpose", "Command or path"], [
        ["Build analytic table", commands[0]],
        ["Regenerate audit and figures", commands[1]],
        ["Regenerate dictionary workbook", commands[2]],
        ["Regenerate this report", commands[3]],
        ["Analysis-ready dataset", "data/processed/hemobp_session_level.csv.gz"],
        ["Machine-readable audit", "artifacts/step2_dataset_audit.json"],
    ], [2.0, 4.7], font_size=8.2)

    doc.add_heading("References", level=1)
    add_body(
        doc,
        "[1] Lin CJ, Chen YY, Pan CF, Wu V, Wu CJ. Dataset supporting blood pressure prediction for the management of chronic hemodialysis. Scientific Data. 2019;6:313. https://doi.org/10.1038/s41597-019-0319-8",
    )
    add_body(
        doc,
        "[2] Chien CY. HEMOBP. Figshare. Version 3. Dataset. 2019. CC BY 4.0. https://doi.org/10.6084/m9.figshare.6260654.v3",
    )

    doc.add_heading("Data Dictionary Conventions", level=1)
    add_body(
        doc,
        "The appendices list every field in the three public source files and every field in the session-level analytic table. Observed ranges describe this specific release; they are not universal physiological limits. Missing percentages use the row count of the relevant table as the denominator. Role and availability state whether a field may be used by a model at the index time. The companion spreadsheet preserves additional provenance and quality-note columns for all 63 variables.",
    )
    conventions = [
        ["Variable", "Exact column name used in the corresponding CSV file"],
        ["Type", "Logical data type after parsing; source files may store values as text"],
        ["Unit", "Physical or coding unit; source-defined index is used where the publication does not document a unit"],
        ["Values or observed range", "Allowed codes for categories or minimum-to-maximum values observed in the audited release"],
        ["Missing", "Number and percentage of rows encoded as null after parsing"],
        ["Role and availability", "Purpose in the analytic design and whether it exists before, at, or after prediction time"],
    ]
    add_table(doc, ["Dictionary element", "Interpretation"], conventions, [2.0, 4.7], font_size=8.4)

    # Landscape appendices for a readable, complete dictionary.
    landscape = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(landscape, landscape=True)

    doc.add_heading("Appendix A Raw Source Data Dictionary", level=1)
    add_body(doc, "The raw dictionary contains 23 variables across Idp, Hemrec D1, and Hemrec VIP. Raw values are never overwritten.")
    for table_name in ("Idp", "Hemrec D1", "Hemrec VIP"):
        if table_name == "Hemrec VIP":
            doc.add_page_break()
        doc.add_heading(table_name, level=2)
        subset = [r for r in raw_rows if r["table"] == table_name]
        formatted = []
        for r in subset:
            formatted.append(
                [
                    r["variable"],
                    r["data_type"],
                    r["unit"],
                    r["allowed_or_observed_values"],
                    missing_text(r),
                    r["definition"],
                    r["analysis_role"],
                ]
            )
        add_table(
            doc,
            ["Variable", "Type", "Unit", "Values or observed range", "Missing", "Definition", "Analytic role"],
            formatted,
            [1.15, 0.72, 0.92, 1.55, 0.78, 2.15, 2.58],
            font_size=7.25,
        )
    add_body(
        doc,
        "Interpretation note. In the source documentation, measuretime is an undocumented measurement index, whereas time is elapsed dialysis time in minutes. They are not treated as interchangeable. Raw minimum and maximum values describe this release; separate index-eligibility checks define the analysis-ready cohort.",
        italic=True,
    )

    doc.add_page_break()
    doc.add_heading("Appendix B Analytic Data Dictionary", level=1)
    add_body(
        doc,
        "The analytic table contains 40 variables. Fields marked audit, eligibility, outcome, or post-index are retained for traceability but are excluded from model predictors.",
    )
    groups = [
        ("Identifiers and Index Documentation", {"pid", "session_date", "index_datetime", "index_minute"}),
        ("Index and Predialysis Predictors", {
            "baseline_sbp", "baseline_dbp", "initial_dialysate_temp_c", "initial_conductivity_ms_cm",
            "initial_uf_l_h", "initial_blood_flow_ml_min", "weightstart", "dryweight", "temperature",
            "gender", "DM", "age_years", "dialysis_vintage_years", "fluid_excess_kg", "fluid_excess_pct",
            "initial_uf_ml_kg_h", "baseline_map", "baseline_pulse_pressure",
        }),
        ("Prior Session Predictors", {"prior_session_idh", "prior_nadir_sbp", "prior_session_count", "prior_idh_rate"}),
        ("Outcome and Post Index Audit", {
            "later_records", "later_distinct_minutes", "last_observed_minute", "nadir_sbp", "mean_later_sbp",
            "idh_absolute", "idh_drop_20", "idh_flythe", "dialysisstart", "dialysisend", "weightend",
            "birthday", "first_dialysis", "session_year",
        }),
    ]
    seen = set()
    for group_name, names in groups:
        subset = [r for r in analytic_rows if r["variable"] in names]
        seen.update(r["variable"] for r in subset)
        doc.add_heading(group_name, level=2)
        formatted = []
        for r in subset:
            role = f'{r["analysis_role"]}; {r["prediction_availability"]}'
            formatted.append(
                [
                    r["variable"],
                    r["data_type"],
                    r["unit"],
                    r["allowed_or_observed_values"],
                    missing_text(r),
                    r["definition"],
                    role,
                ]
            )
        add_table(
            doc,
            ["Variable", "Type", "Unit", "Values or observed range", "Missing", "Definition", "Role and availability"],
            formatted,
            [1.3, 0.7, 0.92, 1.48, 0.75, 2.05, 2.65],
            font_size=7.1,
        )
    unassigned = [r["variable"] for r in analytic_rows if r["variable"] not in seen]
    if unassigned:
        raise RuntimeError(f"Unassigned analytic dictionary variables: {unassigned}")

    configure_headers_and_footers(doc)
    doc.core_properties.title = "DIAL ALERT Data Collection and Understanding"
    doc.core_properties.subject = "Capstone Project Step 2 Dataset Overview and Data Dictionary"
    doc.core_properties.author = "Franklin Guillano"
    doc.core_properties.keywords = "DIAL ALERT, HEMOBP, hemodialysis, data dictionary, capstone"
    doc.core_properties.comments = "DIAL-ALERT academic capstone"
    doc.save(OUT)
    print(f"Wrote {OUT}")
    print(f"Raw variables: {len(raw_rows)}; analytic variables: {len(analytic_rows)}")


if __name__ == "__main__":
    build_report()
