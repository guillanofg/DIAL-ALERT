"""Apply submission wording corrections to existing reports without retraining."""
from pathlib import Path
from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
REPLACEMENTS = {
    "a twelve-slide technical deck for peers and a ten-slide business deck": "a sixteen-slide technical deck for peers and an eleven-slide business deck",
    "0.444 for histogram gradient boosting": "0.446 for histogram gradient boosting",
    "Previous-session nadir SBP produced the largest decrease (0.128), followed by prior IDH rate (0.081), fluid excess percentage (0.015), baseline SBP (0.009), and baseline mean arterial pressure (0.004).": "Previous-session nadir SBP produced the largest decrease (0.076), followed by prior IDH rate (0.059), previous-session IDH (0.017), baseline SBP (0.014), and baseline mean arterial pressure (0.009).",
    "Configurations, split assignments, candidate pipelines, final predictor, threshold, package versions, hashes, and random seeds are saved.": "The final predictor, threshold, configuration and version records are included. Candidate pipelines and patient-level split assignments must be regenerated.",
    "Reproducible patient-disjoint assignment for every session": "Regenerated locally; patient-level assignments are not bundled",
    "The trained pipelines, threshold, configuration, split assignments, metrics, package versions, and integrity hashes are saved.": "The selected predictor, threshold, configuration, metrics and version records are included; other pipelines and patient-level split assignments require regeneration.",

    "follow-up shorter than 120 minutes": "last observed dialysis minute below 120",
    "observation through minute 120": "an observation at or beyond dialysis minute 120 (not 120 minutes after prediction)",
    "observation through at least minute 120": "an observation at or beyond dialysis minute 120",
    "follow-up to at least minute 120": "an observation at or beyond dialysis minute 120 (not 120 minutes after prediction)",
    "Sessions removed for follow-up below 120 minutes": "Sessions with last observed dialysis minute below 120",
    "This rule balanced event ranking with probability reliability and selected the random forest over the less well-calibrated boosted-tree candidate.": "This rule selected random forest using the candidates as fitted. It does not establish superiority after equal calibration of both finalists.",
    "SHAP explanations were generated for the uncalibrated base model because probability calibration does not change the underlying predictor relationships.": "SHAP explanations describe the selected random forest; no post-hoc calibration was applied.",
    "Appendix Rubric Evidence Map": "Appendix Project Evidence Map",
    "Rubric step": "Project step",
}

NOTES = [
    ("Outcome and eligibility", "Predict any SBP below 90 mmHg after the earliest valid active-dialysis observation in minutes 0 to 30. Require index SBP at least 90, two distinct later measurement minutes, and last observed dialysis minute at least 120. The last requirement is anchored to dialysis time, not to prediction time. It is a retrospective eligibility criterion; short or interrupted sessions need separate prospective evaluation."),
    ("Distinct operating strategies", "The fixed threshold is 0.142855878 (rounded to 0.143), chosen by maximum F2 on a patient-disjoint validation decision subset: test sensitivity 66.1%, precision 30.6%, 18.3 reviews and 12.7 false alerts per 100 sessions. Reviewing the highest-risk 20% instead gives 69.1% recall, 29.4% precision and 14.1 false alerts per 100. A prospective capacity policy must define the comparison batch and tie handling before use."),
    ("Model choice and uncertainty", "Random forest was selected from candidates within 0.01 of the best grouped-CV average precision using validation Brier score. No post-hoc calibration was retained. An equally calibrated finalist comparison remains secondary work. The 21,354 test sessions come from only 170 patients. Patient-bootstrap 95% intervals: AP 0.306 to 0.478; ROC AUC 0.821 to 0.876; sensitivity at the fixed threshold 55.0% to 74.5%; recall at 20% capacity 63.2% to 74.4%."),
    ("Proposed future pilot targets", "Provisional targets, not demonstrated prospective results: detect at least 65% of later SBP-below-90 events, generate no more than 15 false alerts per 100 eligible sessions, and achieve median review time at most 2 minutes per displayed alert. Prespecify one operating policy, denominators, uncertainty and safety stopping rules with the local team before evaluation."),
]
FUTURE = [
    ("Optional deployment evidence", "Step 8 includes a local Flask app, a synthetic request, a recorded HTTP demo and instructions. The app and command-line route apply the same prior-session-count transformation. This is a local inference demonstration; no clinical or cloud deployment is established. Step 9 demonstrates saved AI-draft replay with numeric checks. Its published media does not demonstrate live LLM generation."),
    ("Reproducibility status", "A fresh end-to-end reproduction was successfully completed using Python 3.12 on 20 September 2026. The HEMOBP Version 3 source files were downloaded and checksum-verified, the session-level analytical dataset was rebuilt from the raw data, models were retrained and evaluated, EDA and fairness outputs were regenerated, and the automated test suite completed with 15 passed, 0 failed, and 0 skipped. The fresh run was numerically consistent with the locked reference results rather than byte-for-byte identical. Full details and numerical comparisons are recorded in docs/REPRODUCIBILITY_RECORD.md."),
    ("Secondary analyses awaiting execution", "Measure index-to-first-event warning time; compare a baseline-SBP plus prior-hypotension model; assess history ablation and first-observed sessions; compare both finalists with identical calibration splits and methods; and examine extreme UF/fluid-excess values and short-session exclusion. Prespecify analyses on development data, report them as exploratory, and do not replace the locked model based on repeated test-set comparisons."),
    ("Future validation sequence", "Validate the locked model at another center or in a later period, confirm exploratory fairness mitigation on fresh patients, then run silent predictions. Only after data quality and safety gates pass should supervised clinician review assess workload, actions, unintended effects, outcomes and costs. A 90-day schedule is illustrative and cannot establish clinical or financial benefit by itself."),
]

def paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs

def replace_text(p, old, new):
    # Preserve existing run formatting outside the changed span.
    text = p.text
    if old not in text:
        return
    start = text.index(old); end = start + len(old); offset = 0
    for r in p.runs:
        left = offset; right = offset + len(r.text); offset = right
        if right <= start or left >= end:
            continue
        prefix = r.text[:max(0, start-left)]
        suffix = r.text[max(0,end-left):] if right > end else ""
        r.text = prefix + (new if left <= start < right else "") + suffix
    if old in p.text:
        replace_text(p, old, new)

def append_notes(doc):
    if any(p.text == "Submission interpretation and evidence" for p in doc.paragraphs):
        return
    for title, rows in [("Submission interpretation and evidence", NOTES), ("Demonstration status and further research", FUTURE)]:
        doc.add_page_break()
        doc.add_heading(title, level=1)
        for heading, body in rows:
            doc.add_heading(heading, level=2)
            p=doc.add_paragraph(body)
            p.paragraph_format.space_after=Pt(7)
            for r in p.runs:r.font.size=Pt(10)

def main():
    for file in (ROOT / "reports").glob("*.docx"):
        doc=Document(file)
        for p in paragraphs(doc):
            for old,new in REPLACEMENTS.items():replace_text(p,old,new)
        if "Final_Report" in file.name:append_notes(doc)
        doc.save(file)
        print(file.name)

if __name__ == "__main__":main()
