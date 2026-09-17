"""Create the DIAL-ALERT Step 5 Bias and Fairness Analysis report."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from create_eda_report import (
    ARTIFACTS,
    MID_GRAY,
    add_body,
    add_bullets,
    add_code,
    add_figure,
    add_page_break,
    add_table,
    configure_document,
    set_font,
)


ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
REPORTS.mkdir(exist_ok=True)
OUTPUT = REPORTS / "Franklin_Guillano_DIAL_ALERT_Bias_and_Fairness_Analysis.docx"


def percentage(value: float, digits: int = 1) -> str:
    return f"{value:.{digits}%}"


def interval_for(
    intervals: pd.DataFrame,
    attribute: str,
    metric: str,
    digits: int = 3,
) -> str:
    row = intervals[(intervals.attribute == attribute) & (intervals.metric == metric)].iloc[0]
    return f"{row.lower_95:.{digits}f} to {row.upper_95:.{digits}f}"


def make_report() -> Path:
    groups = pd.read_csv(ARTIFACTS / "step5_group_metrics.csv")
    disparities = pd.read_csv(ARTIFACTS / "step5_disparities.csv").set_index("attribute")
    disparity_intervals = pd.read_csv(ARTIFACTS / "step5_disparity_intervals.csv")
    patient_weighted = pd.read_csv(ARTIFACTS / "step5_patient_weighted_disparities.csv").set_index("attribute")
    intersection = pd.read_csv(ARTIFACTS / "step5_intersectional_metrics.csv")
    mitigation = pd.read_csv(ARTIFACTS / "step5_mitigation_comparison.csv").set_index("strategy")
    attributes = pd.read_csv(ARTIFACTS / "step5_attribute_availability.csv")
    robustness = pd.read_csv(ARTIFACTS / "step5_robustness_audit.csv")
    ethical_risks = pd.read_csv(ARTIFACTS / "step5_ethical_risk_register.csv")
    shap_global = pd.read_csv(ARTIFACTS / "step5_shap_global_importance.csv")
    shap_local = pd.read_csv(ARTIFACTS / "step5_shap_local_synthetic.csv")
    synthetic_case = json.loads(
        (ARTIFACTS / "step5_synthetic_case.json").read_text(encoding="utf-8")
    )
    lime_metadata = json.loads((ARTIFACTS / "step5_lime_local_metadata.json").read_text(encoding="utf-8"))
    counterfactual = json.loads((ARTIFACTS / "step5_counterfactual_sex_audit.json").read_text(encoding="utf-8"))
    thresholds = json.loads((ARTIFACTS / "step5_sex_specific_thresholds.json").read_text(encoding="utf-8"))
    audit_summary = json.loads((ARTIFACTS / "step5_audit_summary.json").read_text(encoding="utf-8"))
    final_metrics = json.loads((ARTIFACTS / "final_test_metrics.json").read_text(encoding="utf-8"))[
        "test_metrics_calibrated"
    ]

    sex = groups[groups.attribute == "Recorded sex"].set_index("group")
    age = groups[groups.attribute == "Age group"].set_index("group")
    diabetes = groups[groups.attribute == "Diabetes status"].set_index("group")
    sex_gap = disparities.loc["Recorded sex"]
    age_gap = disparities.loc["Age group"]
    diabetes_gap = disparities.loc["Diabetes status"]
    original = mitigation.loc["Original model"]
    reweighted = mitigation.loc["Sex outcome reweighting"]
    group_threshold = mitigation.loc["Sex specific thresholds"]
    local_first = shap_local.iloc[0]

    doc = Document()
    configure_document(doc)
    doc.sections[0].bottom_margin = Inches(0.82)
    footer = doc.sections[0].footer.paragraphs[0]
    footer.clear()

    # Cover
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_before = Pt(86)
    title.add_run("DIAL ALERT Bias and Fairness Analysis")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("Capstone Project Step 5")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    set_font(
        p.add_run(
            "Ethical review of model explanations subgroup performance uncertainty and mitigation tradeoffs"
        ),
        size=12,
        bold=True,
    )
    p = doc.add_paragraph()
    set_font(p.add_run("Prepared by Franklin Guillano"), size=11)
    p = doc.add_paragraph()
    set_font(p.add_run("September 2026"), size=10.5, color=MID_GRAY)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(96)
    set_font(
        p.add_run(
            "Academic clinical decision support prototype. This audit does not establish clinical benefit, causal fairness, regulatory compliance, or readiness for deployment."
        ),
        size=9.5,
        italic=True,
        color=MID_GRAY,
    )

    add_page_break(doc)
    doc.add_heading("Executive Summary", level=1)
    add_body(
        doc,
        f"The locked Step 4 random forest was audited on {audit_summary['test_sessions']:,} sessions from {audit_summary['test_patients']} patients who were not used for model fitting, hyperparameter selection, calibration choice, or threshold selection. The model achieved test average precision {final_metrics['average_precision']:.3f}, ROC AUC {final_metrics['roc_auc']:.3f}, and Brier score {final_metrics['brier_score']:.3f}. The fairness analysis preserves the validation-selected threshold of {audit_summary['threshold']:.3f} and resamples entire patients when estimating uncertainty.",
    )
    add_body(
        doc,
        f"Point estimates show different fairness concerns across available groups. The recorded-sex disparate impact ratio was {sex_gap.disparate_impact_ratio:.3f} and the equalized odds gap was {sex_gap.equalized_odds_difference:.3f}. Age groups had a disparate impact ratio of {age_gap.disparate_impact_ratio:.3f}, while diabetes status had near-equal alert selection rates but a {diabetes_gap.equal_opportunity_difference:.1%} sensitivity gap. Patient-bootstrap intervals were wide, so the audit identifies monitoring priorities rather than proving discrimination or fairness.",
    )
    add_body(
        doc,
        f"Training-sample reweighting improved the recorded-sex disparate impact ratio from {original.sex_disparate_impact_ratio:.3f} to {reweighted.sex_disparate_impact_ratio:.3f} and reduced the equalized odds gap from {original.sex_equalized_odds_difference:.3f} to {reweighted.sex_equalized_odds_difference:.3f}, with average precision changing from {original.average_precision:.3f} to {reweighted.average_precision:.3f}. A separate sex-specific threshold experiment performed worse on the test set and is rejected. The recommended next step is external and prospective validation with purposeful collection of race, ethnicity, socioeconomic, gender-identity, symptom, intervention, and workflow data.",
    )

    doc.add_heading("Rubric Evidence", level=2)
    rubric_rows = [
        ["Explainability", "Global permutation evidence, aggregate Monte Carlo SHAP, synthetic SHAP and LIME local explanations, and synthetic-profile dependence plots are compared."],
        ["Limitations", "Imbalance, patient and temporal leakage, overfitting, calibration, repeated sessions, external validity, and causal limits are addressed."],
        ["Bias audit", "Recorded sex, age, diabetes, sex-by-age intersections, and patient-equal-weight sensitivity analyses are reported with patient-bootstrap intervals."],
        ["Fairness metrics", "Demographic parity difference, disparate impact ratio, equal opportunity, false-positive-rate, equalized odds, and predictive parity gaps are calculated."],
        ["Mitigation", "Outcome-by-sex reweighting and validation-tuned sex-specific thresholds are tested, including accuracy-fairness tradeoffs and governance risks."],
        ["Ethical AI", "False-negative harm, alert fatigue, automation bias, privacy, accountability, drift, human oversight, and deployment gates are documented."],
    ]
    add_table(doc, ["Requirement", "Evidence"], rubric_rows, widths=[1.45, 5.45], font_size=8.4)

    add_page_break(doc)
    doc.add_heading("1 Audit Scope and Data Sufficiency", level=1)
    add_body(
        doc,
        "The fairness question is whether DIAL-ALERT distributes alerts and errors acceptably across clinically and socially relevant groups. It is not enough to ask whether overall discrimination is good. The audit therefore examines who receives alerts, who is missed, who receives false alerts, how reliable the probabilities are, and whether mitigation transfers from validation patients to new patients.",
    )
    availability_rows = []
    for row in attributes.itertuples(index=False):
        availability_rows.append(
            [row.attribute, row.availability, row.audit_use, row.limitation]
        )
    add_table(
        doc,
        ["Attribute", "Availability", "Audit use", "Important limitation"],
        availability_rows,
        widths=[1.25, 1.05, 1.7, 2.9],
        font_size=7.8,
        vertical_margin=65,
    )
    add_body(
        doc,
        "The source field is named gender but contains only F and M. This report describes it as recorded sex because it cannot represent gender identity. Race, ethnicity, and socioeconomic status are absent. Their fairness cannot be inferred from diabetes, dialysis vintage, geography, or another convenient proxy. Missing protected-attribute data are a substantive limitation and a requirement for future data collection.",
        bold_lead="The source field is named gender but contains only F and M.",
    )

    doc.add_heading("Audit Population", level=2)
    audit_rows = [
        ["Test observations", f"{audit_summary['test_sessions']:,} sessions"],
        ["Independent test units", f"{audit_summary['test_patients']} patients"],
        ["Patient overlap", "Zero across training validation and test partitions"],
        ["Operating rule", f"Global probability threshold {audit_summary['threshold']:.3f} selected on validation patients"],
        ["Uncertainty", f"{audit_summary['bootstrap_iterations']} bootstrap resamples of whole patients"],
    ]
    add_table(doc, ["Element", "Specification"], audit_rows, widths=[1.8, 5.1], font_size=8.9)

    add_page_break(doc)
    doc.add_heading("2 Ethical Framing and Harm Model", level=1)
    add_body(
        doc,
        "DIAL-ALERT predicts a later systolic blood pressure below 90 mmHg. The target is an auditable blood-pressure event, not a complete clinical diagnosis of intradialytic hypotension because symptoms, treatments, and clinician assessments are unavailable. A risk score should therefore support attention and preparation, not autonomously diagnose, prescribe fluid, change ultrafiltration, or terminate dialysis.",
    )
    harm_rows = [
        ["False negative", "An at-risk session is not flagged", "Delayed recognition or preparation", "Sensitivity and missed-event review"],
        ["False positive", "A low-risk session is flagged", "Alarm fatigue unnecessary review or intervention", "Precision false-positive rate and alert capacity"],
        ["Unequal false negatives", "One group is missed more often", "Unequal safety benefit", "Equal opportunity difference"],
        ["Unequal false positives", "One group receives more false alerts", "Unequal burden and possible overtreatment", "False-positive-rate difference"],
        ["Unequal selection", "Alert exposure differs regardless of outcome", "Different access to review or intervention", "Demographic parity and disparate impact"],
        ["Miscalibration", "The same score means different risk across groups", "Misleading urgency and threshold decisions", "Brier score and calibration error"],
    ]
    add_table(doc, ["Concern", "Meaning", "Potential harm", "Audit measure"], harm_rows, widths=[1.35, 2.0, 2.0, 1.55], font_size=7.9)
    add_body(
        doc,
        "No single fairness metric is sufficient. Different outcome prevalence and an imperfect classifier make some fairness goals mutually difficult to satisfy. Metric choice must follow a documented clinical harm model, workflow capacity, and community governance rather than a numerical pass-fail rule alone.",
    )

    add_page_break(doc)
    doc.add_heading("3 Explanation Framework", level=1)
    add_body(
        doc,
        "Four complementary explanation methods were used because each answers a different question. Agreement across methods strengthens confidence in a descriptive model interpretation, while disagreement or low local fidelity is itself evidence that an explanation should not be trusted quantitatively.",
    )
    explanation_rows = [
        ["Permutation importance", "What information uniquely improves test-set ranking", "Can understate correlated or redundant predictors"],
        ["Approximate SHAP", "How features distribute a prediction relative to background sessions", "Monte Carlo and background dependent; correlated features complicate attribution"],
        ["LIME-style surrogate", "Which synthetic scenario features influence one nearby linear approximation", "Only trustworthy when local fidelity is adequate"],
        ["PDP", "Average prediction change when one feature is varied", "Can create unrealistic combinations and hides heterogeneity"],
        ["Synthetic-profile ICE", "How the same feature change affects predefined synthetic profiles", "Illustrates model behavior, not observed patients or treatment effects"],
    ]
    add_table(doc, ["Method", "Question answered", "Main caution"], explanation_rows, widths=[1.5, 2.65, 2.75], font_size=8.3)
    add_body(
        doc,
        "The SHAP implementation uses random background sessions and feature-order permutations to estimate interventional Shapley contributions. Only aggregate global importance is published. The local path ends at a manually constructed synthetic scenario, so no local explanation represents an actual HEMOBP session. Feature substitution can still form clinically implausible combinations. These values explain the fitted model, not physiology or causation.",
    )

    add_page_break(doc)
    doc.add_heading("4 Global Model Explanation", level=1)
    add_figure(
        doc,
        ARTIFACTS / "step5_shap_summary.png",
        "Figure 1. Aggregate mean absolute approximate SHAP contributions from an outcome-balanced analysis sample.",
        width=6.5,
        alt_text="Aggregate bar chart showing prior-session nadir blood pressure and prior IDH rate as dominant global contributors.",
    )
    top_features = shap_global.head(7)
    shap_rows = [
        [
            row.feature.replace("_", " ").title(),
            f"{row.mean_absolute_shap:.3f}",
            f"{row.mean_shap:+.3f}",
        ]
        for row in top_features.itertuples(index=False)
    ]
    add_table(doc, ["Feature", "Mean absolute contribution", "Mean signed contribution"], shap_rows, widths=[3.1, 1.9, 1.9], font_size=8.7)
    add_body(
        doc,
        "Lower prior-session nadir SBP and higher prior IDH rate dominate risk predictions, followed by fluid excess and prior-session IDH. The same variables lead the independent permutation analysis. Recorded sex and age contribute little unique global information relative to the longitudinal haemodynamic features, but low feature importance does not rule out subgroup disparity through correlated predictors, threshold effects, or different baseline risk.",
    )

    doc.add_heading("5 Synthetic Local Explanations", level=1)
    add_body(
        doc,
        f"A predefined synthetic teaching scenario had a predicted probability of {local_first.prediction:.3f}, compared with a background expectation of {local_first.baseline_prediction:.3f}. The approximate SHAP contributions sum to the model prediction with local accuracy error {abs(local_first.local_accuracy_error):.3g}. This scenario is explicitly marked `{synthetic_case['scenario_type']}` and is not copied from a patient record. Prior IDH rate, prior-session nadir SBP, fluid excess, and a prior-session event produced the largest upward contributions.",
    )
    add_figure(
        doc,
        ARTIFACTS / "step5_shap_local_synthetic.png",
        "Figure 2. Approximate SHAP contributions for a synthetic high-risk teaching scenario.",
        width=6.65,
        alt_text="Horizontal bars show prior IDH rate, previous nadir SBP, and fluid excess raising predicted risk for a synthetic scenario.",
    )
    add_body(
        doc,
        "The explanation is clinically plausible but not prescriptive. For example, the positive contribution from fluid excess does not prove that changing fluid removal would reduce risk. An actionable recommendation requires causal and prospective evidence that is outside this retrospective prediction study.",
    )

    add_page_break(doc)
    doc.add_heading("6 LIME Style Local Surrogate", level=1)
    add_figure(
        doc,
        ARTIFACTS / "step5_lime_local_synthetic.png",
        "Figure 3. Weighted LIME-style linear surrogate for the same synthetic scenario.",
        width=6.6,
        alt_text="Local surrogate coefficients broadly agree with SHAP but show only moderate quantitative fidelity.",
    )
    lime_rows = [
        ["Perturbation samples", f"{lime_metadata['samples']:,}"],
        ["Weighted local R squared", f"{lime_metadata['weighted_r2']:.3f}"],
        ["Random forest probability", f"{lime_metadata['model_prediction']:.3f}"],
        ["Surrogate probability", f"{lime_metadata['surrogate_prediction']:.3f}"],
        ["Absolute local error", f"{lime_metadata['absolute_prediction_error']:.3f}"],
    ]
    add_table(doc, ["Diagnostic", "Result"], lime_rows, widths=[3.3, 3.0], font_size=9.0)
    add_body(
        doc,
        f"The surrogate agrees qualitatively with SHAP on the dominant features, but its weighted R squared of {lime_metadata['weighted_r2']:.3f} and absolute probability error of {lime_metadata['absolute_prediction_error']:.3f} are not strong enough for quantitative reliance. The LIME-style result is therefore reported as corroborating evidence with a fidelity warning, not as a substitute for the random forest. This limitation illustrates why every local explanation requires its own quality check.",
    )

    add_page_break(doc)
    doc.add_heading("7 Synthetic Profile Dependence", level=1)
    add_figure(
        doc,
        ARTIFACTS / "step5_pdp_ice.png",
        "Figure 4. Dependence averages and conditional-effect curves across five synthetic reference profiles.",
        width=6.9,
        alt_text="Average risk falls as prior nadir SBP rises and increases as prior IDH rate or fluid excess rises across synthetic profiles.",
    )
    add_body(
        doc,
        "Average predicted risk falls sharply as prior-session nadir SBP rises from very low values. It increases with prior IDH rate and, more gradually, with fluid excess percentage. The spread of the synthetic-profile curves shows model heterogeneity without displaying any observed patient trajectory.",
    )
    add_body(
        doc,
        "The curves vary one feature while leaving the other synthetic-profile values fixed. Because blood pressure, dry weight, fluid excess, and ultrafiltration measures are correlated, some combinations may be physiologically uncommon. These plots should be used to inspect model behavior, not to infer a safe treatment target or causal dose-response relationship.",
    )

    add_page_break(doc)
    doc.add_heading("8 Imbalance Leakage and Overfitting", level=1)
    robustness_rows = []
    for row in robustness.itertuples(index=False):
        evidence = row.evidence
        if row.risk == "Patient leakage":
            evidence = "Zero patient overlap across training, validation, and test partitions"
        robustness_rows.append([row.risk, evidence, row.assessment, row.response])
    add_table(
        doc,
        ["Risk", "Observed evidence", "Assessment", "Required response"],
        robustness_rows,
        widths=[1.25, 2.65, 1.25, 1.75],
        font_size=7.4,
        vertical_margin=55,
    )
    add_body(
        doc,
        "Class imbalance is handled analytically rather than erased. Average precision, sensitivity, precision, calibration, and alert workload remain primary. Accuracy is not used to justify the model because predicting no event would already appear accurate for an outcome with approximately 8.5% prevalence.",
    )
    add_body(
        doc,
        "Patient leakage is controlled by disjoint patient partitions and patient-grouped cross-validation. Temporal leakage is controlled by using only index-time or earlier-session predictors. Residual risks include patient-date linkage ambiguity, informative monitoring frequency, and the possibility that operational data fields would be recorded differently after deployment.",
    )
    add_body(
        doc,
        "Overfitting remains plausible because the test average precision of 0.395 is lower than the validation value of 0.470 and because only 170 independent patients are in the test partition. Patient-bootstrap intervals, a locked test, and simpler comparators reduce optimism but do not replace external temporal and geographic validation.",
    )

    add_page_break(doc)
    doc.add_heading("9 Fairness Metrics", level=1)
    metric_rows = [
        ["Demographic parity difference", "Maximum minus minimum alert-selection rate", "Who receives an alert regardless of outcome"],
        ["Disparate impact ratio", "Minimum divided by maximum selection rate", "Relative alert exposure; 0.80 is shown only as a screening heuristic"],
        ["Equal opportunity difference", "Maximum minus minimum sensitivity", "Who receives the safety benefit among sessions with events"],
        ["False-positive-rate difference", "Maximum minus minimum FPR", "Who bears unnecessary alerts among sessions without events"],
        ["Equalized odds difference", "Larger of sensitivity and FPR gaps", "Worst conditional error-rate disparity"],
        ["Predictive parity difference", "Maximum minus minimum precision", "Whether an alert has similar meaning across groups"],
        ["Brier-score difference", "Maximum minus minimum probability error", "Whether probability quality differs across groups"],
    ]
    add_table(doc, ["Metric", "Definition", "Ethical interpretation"], metric_rows, widths=[1.8, 2.45, 2.65], font_size=8.0)
    add_body(
        doc,
        "The 0.80 ratio is a descriptive screening heuristic borrowed from another policy context, not a legal, regulatory, statistical, or clinical fairness certificate for DIAL-ALERT. Direction, uncertainty, prevalence, error severity, and workflow consequences must be reviewed together.",
    )
    add_figure(
        doc,
        ARTIFACTS / "step5_disparity_summary.png",
        "Figure 5. Disparate impact and equalized odds summarize different dimensions of subgroup disparity.",
        width=6.55,
        alt_text="Recorded sex and age have selection-rate ratios below 0.80, while diabetes shows a larger sensitivity than selection-rate disparity.",
    )

    add_page_break(doc)
    doc.add_heading("10 Subgroup Audit Results", level=1)
    add_figure(
        doc,
        ARTIFACTS / "step5_fairness_overview.png",
        "Figure 6. Alert selection sensitivity and false-positive rates by recorded sex age group and diabetes status.",
        width=6.75,
        alt_text="Subgroup bars show higher alert and false-positive rates among female sessions and age 55 to 64, and lower sensitivity among sessions with diabetes.",
    )
    subgroup_rows = [
        ["Recorded sex F", f"{int(sex.loc['F'].patients)}", percentage(sex.loc["F"].prevalence), percentage(sex.loc["F"].selection_rate), percentage(sex.loc["F"].sensitivity), percentage(sex.loc["F"].false_positive_rate), percentage(sex.loc["F"].precision)],
        ["Recorded sex M", f"{int(sex.loc['M'].patients)}", percentage(sex.loc["M"].prevalence), percentage(sex.loc["M"].selection_rate), percentage(sex.loc["M"].sensitivity), percentage(sex.loc["M"].false_positive_rate), percentage(sex.loc["M"].precision)],
        ["Age under 55", f"{int(age.loc['Under 55'].patients)}", percentage(age.loc["Under 55"].prevalence), percentage(age.loc["Under 55"].selection_rate), percentage(age.loc["Under 55"].sensitivity), percentage(age.loc["Under 55"].false_positive_rate), percentage(age.loc["Under 55"].precision)],
        ["Age 55 to 64", f"{int(age.loc['55 to 64'].patients)}", percentage(age.loc["55 to 64"].prevalence), percentage(age.loc["55 to 64"].selection_rate), percentage(age.loc["55 to 64"].sensitivity), percentage(age.loc["55 to 64"].false_positive_rate), percentage(age.loc["55 to 64"].precision)],
        ["Age 65 to 74", f"{int(age.loc['65 to 74'].patients)}", percentage(age.loc["65 to 74"].prevalence), percentage(age.loc["65 to 74"].selection_rate), percentage(age.loc["65 to 74"].sensitivity), percentage(age.loc["65 to 74"].false_positive_rate), percentage(age.loc["65 to 74"].precision)],
        ["Age 75 and older", f"{int(age.loc['75 and older'].patients)}", percentage(age.loc["75 and older"].prevalence), percentage(age.loc["75 and older"].selection_rate), percentage(age.loc["75 and older"].sensitivity), percentage(age.loc["75 and older"].false_positive_rate), percentage(age.loc["75 and older"].precision)],
        ["Diabetes", f"{int(diabetes.loc['Diabetes'].patients)}", percentage(diabetes.loc["Diabetes"].prevalence), percentage(diabetes.loc["Diabetes"].selection_rate), percentage(diabetes.loc["Diabetes"].sensitivity), percentage(diabetes.loc["Diabetes"].false_positive_rate), percentage(diabetes.loc["Diabetes"].precision)],
        ["No diabetes", f"{int(diabetes.loc['No diabetes'].patients)}", percentage(diabetes.loc["No diabetes"].prevalence), percentage(diabetes.loc["No diabetes"].selection_rate), percentage(diabetes.loc["No diabetes"].sensitivity), percentage(diabetes.loc["No diabetes"].false_positive_rate), percentage(diabetes.loc["No diabetes"].precision)],
    ]
    add_table(
        doc,
        ["Group", "Patients", "Prev", "Alert", "Sens", "FPR", "PPV"],
        subgroup_rows,
        widths=[1.75, 0.75, 0.85, 0.85, 0.85, 0.85, 0.85],
        font_size=7.8,
        vertical_margin=55,
    )

    add_page_break(doc)
    doc.add_heading("11 Interpretation of Subgroup Findings", level=1)
    add_body(
        doc,
        f"Recorded sex. Female sessions had higher outcome prevalence ({sex.loc['F'].prevalence:.1%} versus {sex.loc['M'].prevalence:.1%}), higher alert selection ({sex.loc['F'].selection_rate:.1%} versus {sex.loc['M'].selection_rate:.1%}), higher sensitivity ({sex.loc['F'].sensitivity:.1%} versus {sex.loc['M'].sensitivity:.1%}), and higher false-positive rate ({sex.loc['F'].false_positive_rate:.1%} versus {sex.loc['M'].false_positive_rate:.1%}). The disparate impact ratio was {sex_gap.disparate_impact_ratio:.3f}, with a patient-bootstrap interval of {interval_for(disparity_intervals, 'Recorded sex', 'disparate_impact_ratio')}. The equalized odds gap was {sex_gap.equalized_odds_difference:.3f}, with interval {interval_for(disparity_intervals, 'Recorded sex', 'equalized_odds_difference')}. These wide intervals preclude a precise fairness conclusion.",
        bold_lead="Recorded sex.",
    )
    add_body(
        doc,
        f"Age. Alert selection ranged from {age.selection_rate.min():.1%} to {age.selection_rate.max():.1%}, producing a disparate impact ratio of {age_gap.disparate_impact_ratio:.3f}. Sensitivity ranged from {age.sensitivity.min():.1%} to {age.sensitivity.max():.1%}. Age is both a predictor and a clinically relevant protected characteristic, so an apparent disparity can reflect prevalence, model behavior, threshold choice, or all three. It cannot be labelled unfair or acceptable without a clinical harm analysis.",
        bold_lead="Age.",
    )
    add_body(
        doc,
        f"Diabetes. Alert selection was almost identical ({diabetes.loc['Diabetes'].selection_rate:.1%} versus {diabetes.loc['No diabetes'].selection_rate:.1%}), giving a disparate impact ratio of {diabetes_gap.disparate_impact_ratio:.3f}. Nevertheless, sensitivity was lower with diabetes ({diabetes.loc['Diabetes'].sensitivity:.1%} versus {diabetes.loc['No diabetes'].sensitivity:.1%}) and average precision was also lower ({diabetes.loc['Diabetes'].average_precision:.3f} versus {diabetes.loc['No diabetes'].average_precision:.3f}). Selection parity alone would therefore miss an important performance gap.",
        bold_lead="Diabetes.",
    )
    add_body(
        doc,
        f"Patient weighting and intersections. Giving each patient equal total weight changed several point estimates: recorded-sex disparate impact became {patient_weighted.loc['Recorded sex'].disparate_impact_ratio:.3f}, and the age equalized odds gap became {patient_weighted.loc['Age group'].equalized_odds_difference:.3f}. In the sex-by-age analysis, subgroup patient counts ranged from {int(intersection.patients.min())} to {int(intersection.patients.max())}, and sensitivity ranged from {intersection.sensitivity.min():.1%} to {intersection.sensitivity.max():.1%}. These swings show that frequent attenders and small intersections materially influence results; they justify continued monitoring rather than subgroup-specific deployment claims.",
        bold_lead="Patient weighting and intersections.",
    )

    add_page_break(doc)
    doc.add_heading("12 Direct Dependence on Recorded Sex", level=1)
    add_body(
        doc,
        f"A limited counterfactual diagnostic swapped F and M while holding all other predictors fixed. The mean absolute probability change was {counterfactual['mean_absolute_probability_change']:.3f}, the median was {counterfactual['median_absolute_probability_change']:.3f}, and {counterfactual['alert_class_changed_percent']:.2f}% of session classifications changed at the locked threshold. The largest individual probability change was {counterfactual['maximum_absolute_probability_change']:.3f}.",
    )
    add_body(
        doc,
        "This result suggests that the recorded-sex field has modest direct influence on most predictions. It does not show that the model is free from sex-related bias. Other predictors can encode correlated differences, and changing recorded sex while holding physiology fixed is not a biologically or socially valid causal intervention. The diagnostic must be read together with the observed subgroup error rates.",
    )
    counter_rows = [
        ["Mean absolute probability change", f"{counterfactual['mean_absolute_probability_change']:.4f}"],
        ["Median absolute probability change", f"{counterfactual['median_absolute_probability_change']:.4f}"],
        ["Maximum absolute probability change", f"{counterfactual['maximum_absolute_probability_change']:.4f}"],
        ["Alert classifications changed", f"{counterfactual['alert_class_changed_percent']:.2f}%"],
    ]
    add_table(doc, ["Diagnostic", "Result"], counter_rows, widths=[4.0, 2.2], font_size=9.2)

    add_page_break(doc)
    doc.add_heading("13 Mitigation Experiments", level=1)
    add_figure(
        doc,
        ARTIFACTS / "step5_mitigation_tradeoff.png",
        "Figure 7. Overall performance and recorded-sex fairness metrics for two experimental mitigations.",
        width=6.8,
        alt_text="Reweighting modestly improves recorded-sex fairness with stable average precision, while sex-specific thresholds worsen test fairness.",
    )
    mitigation_rows = [
        ["Original model", f"{original.average_precision:.3f}", f"{original.sensitivity:.1%}", f"{original.precision:.1%}", f"{original.sex_disparate_impact_ratio:.3f}", f"{original.sex_equalized_odds_difference:.3f}"],
        ["Sex-outcome reweighting", f"{reweighted.average_precision:.3f}", f"{reweighted.sensitivity:.1%}", f"{reweighted.precision:.1%}", f"{reweighted.sex_disparate_impact_ratio:.3f}", f"{reweighted.sex_equalized_odds_difference:.3f}"],
        ["Sex-specific thresholds", f"{group_threshold.average_precision:.3f}", f"{group_threshold.sensitivity:.1%}", f"{group_threshold.precision:.1%}", f"{group_threshold.sex_disparate_impact_ratio:.3f}", f"{group_threshold.sex_equalized_odds_difference:.3f}"],
    ]
    add_table(doc, ["Strategy", "AP", "Sens", "PPV", "DI ratio", "EO gap"], mitigation_rows, widths=[2.45, 0.75, 0.75, 0.75, 1.0, 0.9], font_size=8.5)
    add_body(
        doc,
        f"Reweighting. Training rows were weighted so recorded-sex and outcome combinations contributed according to their marginal frequencies. This improved both recorded-sex fairness summaries with negligible overall performance change, but the disparate impact ratio remained below 0.80. Reweighting is the most promising tested option and should be re-evaluated with external data and broader attributes.",
        bold_lead="Reweighting.",
    )
    add_body(
        doc,
        f"Sex-specific thresholds. Validation-tuned thresholds were {thresholds['thresholds']['F']:.3f} for F and {thresholds['thresholds']['M']:.3f} for M, targeting equal sensitivity on validation patients. On the test set, the equalized odds gap worsened from {original.sex_equalized_odds_difference:.3f} to {group_threshold.sex_equalized_odds_difference:.3f}, precision fell from {original.precision:.1%} to {group_threshold.precision:.1%}, and the disparate impact ratio fell to {group_threshold.sex_disparate_impact_ratio:.3f}. The approach failed to generalize and is rejected. It also requires a protected attribute at inference, which raises additional legal and ethical governance questions.",
        bold_lead="Sex-specific thresholds.",
    )

    add_page_break(doc)
    doc.add_heading("14 Mitigation and Governance Plan", level=1)
    mitigation_plan = [
        ["Before modelling", "Collect race ethnicity gender identity socioeconomic indicators symptoms treatments and site variables with consent and governance", "Enables missing audits and improves target validity"],
        ["Pre-processing", "Evaluate patient-balanced and sex-outcome reweighting using validation data only", "Reduces overrepresentation and selected error gaps"],
        ["In-processing", "Constrain sensitivity and FPR gaps during model selection when sample size supports it", "Makes fairness part of optimisation rather than a late check"],
        ["Post-processing", "Use a single capacity-aware threshold by default; test group thresholds only under explicit governance", "Avoids unstable or opaque differential treatment"],
        ["Human factors", "Display probability uncertainty key drivers data quality and a documented override path", "Reduces automation bias and supports accountability"],
        ["Monitoring", "Track overall and subgroup calibration sensitivity precision FPR selection rate and alert burden", "Detects drift and emergent disparity"],
        ["Rollback", "Suspend alerts when data drift calibration failure or subgroup harm thresholds are breached", "Limits harm while the model is reviewed"],
    ]
    add_table(doc, ["Stage", "Action", "Purpose"], mitigation_plan, widths=[1.25, 3.75, 1.9], font_size=7.8, vertical_margin=60)
    add_body(
        doc,
        "A mitigation should not be accepted merely because one metric improves. The acceptance decision should require stable external results, preserved clinical utility, acceptable alert burden, transparent governance, and no material worsening in another subgroup or error type. DIAL-ALERT should remain decision support with accountable clinician override.",
    )

    add_page_break(doc)
    doc.add_heading("15 Ethical Risk Register", level=1)
    risk_rows = [
        [row.hazard, row.stakeholder, row.severity, row.mitigation, row.residual_risk]
        for row in ethical_risks.itertuples(index=False)
    ]
    add_table(
        doc,
        ["Hazard", "Affected party", "Severity", "Control", "Residual risk"],
        risk_rows,
        widths=[1.35, 1.25, 0.75, 2.2, 1.35],
        font_size=7.4,
        vertical_margin=55,
    )
    doc.add_heading("Deployment Gates", level=2)
    add_bullets(
        doc,
        [
            "External temporal and geographic validation with patient-level uncertainty.",
            "Prospective silent-mode evaluation before any alert is shown to clinicians.",
            "Predefined subgroup sample-size and performance thresholds, including currently unavailable groups.",
            "Human-factors testing for alert comprehension, automation bias, alarm fatigue, and override documentation.",
            "Privacy, security, access-control, audit-log, incident-response, rollback, and model-change procedures.",
            "Clinical governance approval that specifies intended use, excluded uses, accountable owner, monitoring frequency, and stopping rules.",
        ],
    )
    add_body(
        doc,
        "Until these gates are met, DIAL-ALERT should be presented only as a retrospective academic prototype. It must not determine treatment, deny care, or substitute for clinical assessment.",
    )

    add_page_break(doc)
    doc.add_heading("16 Conclusions", level=1)
    add_body(
        doc,
        "The audit identifies a clinically coherent but imperfect model. Aggregate global evidence and synthetic local explanations consistently emphasize prior haemodynamic instability and fluid status. Synthetic-profile dependence curves show nonlinear and heterogeneous behavior. The LIME-style surrogate supports the broad feature ranking but has only moderate local fidelity, so it should not be treated as a precise explanation.",
    )
    add_body(
        doc,
        "Fairness conclusions depend on the metric. Recorded sex and age show notable alert-selection differences; diabetes has near-parity in selection but lower sensitivity and discrimination. Patient-bootstrap intervals are wide, patient-equal weighting changes point estimates, and small intersections are unstable. Race, ethnicity, socioeconomic status, and gender identity remain unaudited because the data do not contain them.",
    )
    add_body(
        doc,
        "Sex-outcome reweighting is the only tested mitigation that improved both recorded-sex disparity summaries while preserving overall performance, although it did not eliminate disparity. Sex-specific thresholds failed on the held-out test patients and should not be used. The defensible decision is to continue research with broader data, external validation, patient and staff participation, and explicit deployment controls rather than claim fairness or clinical readiness.",
    )
    doc.add_heading("Reproducibility", level=2)
    add_code(
        doc,
        [
            "python src/audit_bias_fairness.py",
            "python src/create_bias_fairness_report.py",
        ],
    )

    doc.add_heading("References", level=1)
    references = [
        "1. Lin CJ, Chen YY, Pan CF, Wu VC, Wu CJ. Dataset supporting blood pressure prediction for the management of chronic hemodialysis. Scientific Data. 2019;6:313. https://doi.org/10.1038/s41597-019-0319-8",
        "2. Lundberg SM, Lee SI. A unified approach to interpreting model predictions. Advances in Neural Information Processing Systems. 2017;30. https://papers.neurips.cc/paper/7062-a-unified-approach-to-interpreting-model-predictions",
        "3. Ribeiro MT, Singh S, Guestrin C. Why should I trust you Explaining the predictions of any classifier. Proceedings of KDD. 2016:1135-1144. https://doi.org/10.1145/2939672.2939778",
        "4. Goldstein A, Kapelner A, Bleich J, Pitkin E. Peeking inside the black box Visualizing statistical learning with plots of individual conditional expectation. Journal of Computational and Graphical Statistics. 2015;24:44-65. https://doi.org/10.1080/10618600.2014.907095",
        "5. Hardt M, Price E, Srebro N. Equality of opportunity in supervised learning. Advances in Neural Information Processing Systems. 2016;29. https://arxiv.org/abs/1610.02413",
        "6. Chouldechova A. Fair prediction with disparate impact A study of bias in recidivism prediction instruments. Big Data. 2017;5:153-163. https://doi.org/10.1089/big.2016.0047",
        "7. World Health Organization. Ethics and governance of artificial intelligence for health. 2021. https://www.who.int/publications/i/item/9789240029200",
        "8. Collins GS and colleagues. TRIPOD plus AI statement Updated guidance for reporting clinical prediction models that use regression or machine learning methods. BMJ. 2024;385:e078378. https://doi.org/10.1136/bmj-2023-078378",
    ]
    for reference in references:
        add_body(doc, reference)

    doc.core_properties.title = "DIAL ALERT Bias and Fairness Analysis"
    doc.core_properties.subject = "Capstone Project Step 5"
    doc.core_properties.author = "Franklin Guillano"
    doc.core_properties.keywords = "DIAL-ALERT, ethical AI, fairness, bias, SHAP, LIME, PDP, ICE"
    doc.core_properties.comments = "DIAL-ALERT academic capstone"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(make_report())
