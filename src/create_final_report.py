"""Create the consolidated DIAL-ALERT capstone final report."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from create_eda_report import (
    BLACK,
    MID_GRAY,
    NAVY,
    add_body,
    add_bullets,
    add_figure,
    add_page_break,
    add_table,
    configure_document,
    set_font,
)


ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS = ROOT / "artifacts"
REPORTS = ROOT / "reports"
REPORTS.mkdir(exist_ok=True)
OUTPUT = REPORTS / "Franklin_Guillano_DIAL_ALERT_Final_Report.docx"

TEAL = "0E7C7B"
PALE_TEAL = "DCEFEB"
CORAL = "E76F51"


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def replace_footer(doc: Document) -> None:
    footer = doc.sections[0].footer
    paragraph = footer.paragraphs[0]
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("DIAL ALERT   |   Capstone Final Report   |   ")
    set_font(run, size=8.5, color=MID_GRAY)
    page_run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    page_run._r.extend([begin, instruction, end])
    set_font(page_run, size=8.5, color=MID_GRAY)


def add_title_page(doc: Document, cohort: dict, metrics: dict) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(24)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("DIAL ALERT Capstone Final Report")
    set_font(run, size=29, bold=True, color=BLACK)

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(
        "Early prediction of blood pressure defined intradialytic hypotension"
    )
    set_font(run, size=15, color=MID_GRAY)

    author = doc.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.paragraph_format.space_before = Pt(18)
    run = author.add_run("Franklin B. Guillano")
    set_font(run, size=12, bold=True)
    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = date.add_run("September 2026")
    set_font(run, size=10.5, color=MID_GRAY)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    add_table(
        doc,
        ["Analytic cohort", "Primary event rate", "Test ROC AUC", "Test average precision"],
        [[
            f"{cohort['final_eligible_sessions']:,} sessions",
            f"{100 * cohort['primary_idh_prevalence']:.2f}%",
            f"{metrics['roc_auc']:.3f}",
            f"{metrics['average_precision']:.3f}",
        ]],
        widths=[1.75, 1.55, 1.45, 1.75],
        font_size=9.0,
        vertical_margin=150,
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(22)
    p.paragraph_format.left_indent = Inches(0.55)
    p.paragraph_format.right_indent = Inches(0.55)
    run = p.add_run(
        "Academic clinical decision support prototype. The model is not a diagnostic device and is not ready for clinical deployment."
    )
    set_font(run, size=10.5, bold=True, color=CORAL)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.left_indent = Inches(0.45)
    p.paragraph_format.right_indent = Inches(0.45)
    run = p.add_run(
        "This report consolidates problem framing, data understanding, preprocessing, exploratory analysis, feature engineering, model comparison, explainability, fairness auditing, business translation, and reproducibility."
    )
    set_font(run, size=10.5)
    add_page_break(doc)


def add_report_map(doc: Document) -> None:
    doc.add_heading("Report Structure", level=1)
    add_body(
        doc,
        "The report follows the machine-learning lifecycle and the seven graded capstone steps. Each section distinguishes observed evidence from interpretation and proposed future work."
    )
    add_table(
        doc,
        ["Section", "Purpose"],
        [
            ["Executive Summary", "Main result, operational meaning, and recommendation"],
            ["1 Problem Understanding and Framing", "Clinical context, task type, metrics, and business KPIs"],
            ["2 Data Collection and Understanding", "Provenance, cohort, variables, quality, and license"],
            ["3 Preprocessing EDA and Feature Engineering", "Cleaning, leakage control, relationships, selection, and PCA"],
            ["4 Model Implementation and Comparison", "Algorithms, grouped validation, locked test performance, and capacity"],
            ["5 Critical Thinking Ethical AI and Bias Auditing", "Explanations, limitations, fairness metrics, and mitigations"],
            ["6 Final Communication and Business Translation", "Workflow, ROI framework, risks, and pilot design"],
            ["7 Reproducibility and GitHub Repository", "Environment, artifacts, tests, CI, and repository structure"],
            ["8 Conclusions and Next Steps", "Decision, validation gates, and research priorities"],
        ],
        widths=[2.55, 4.15],
        font_size=8.7,
    )
    add_page_break(doc)


def add_executive_summary(doc: Document, cohort: dict, metrics: dict) -> None:
    doc.add_heading("Executive Summary", level=1)
    add_body(
        doc,
        "DIAL-ALERT tests whether information available early in a haemodialysis session can identify sessions that later develop a systolic blood-pressure reading below 90 mmHg. The project uses HEMOBP Version 3, a public longitudinal dataset released under CC BY 4.0. Deterministic linkage and eligibility rules produced 106,758 sessions from 830 patients, with 9,067 primary events. [1,2]"
    )
    add_body(
        doc,
        "Seven supervised modelling configurations were compared using patient-grouped cross-validation and patient-disjoint validation. Random forest was selected under a locked rule that prioritized grouped-CV average precision and then validation Brier score when models were within 0.01 of the best result. The untouched test set contained 21,354 sessions from 170 patients."
    )
    add_body(
        doc,
        f"On the locked test set, the selected model achieved average precision {metrics['average_precision']:.3f}, ROC AUC {metrics['roc_auc']:.3f}, and Brier score {metrics['brier_score']:.3f}. At the validation-selected threshold of {metrics['threshold']:.3f}, sensitivity was {100 * metrics['sensitivity']:.1f}%, specificity {100 * metrics['specificity']:.1f}%, and precision {100 * metrics['precision']:.1f}%. At a 20% alert capacity, the model captured {100 * metrics['recall_at_capacity']:.1f}% of observed events with {100 * metrics['precision_at_capacity']:.1f}% precision and {metrics['lift_at_capacity']:.2f}-fold lift over prevalence."
    )
    add_body(
        doc,
        "The model concentrates retrospective risk but does not prove that an alert prevents an event or produces a financial return. The appropriate decision is to continue development through external validation and a governed prospective silent-mode pilot. Clinical use should remain prohibited until performance, calibration, workload, subgroup outcomes, and human factors meet prespecified acceptance criteria."
    )
    add_table(
        doc,
        ["Decision area", "Conclusion"],
        [
            ["Technical feasibility", "Supported for continued validation; discrimination and risk concentration are promising"],
            ["Clinical effectiveness", "Unknown; no alert-triggered intervention was tested"],
            ["Operational feasibility", "Requires capacity testing because false alerts remain substantial"],
            ["Fairness", "Partially auditable; recorded-sex and age disparities require monitoring, while key attributes are absent"],
            ["Deployment readiness", "Not ready; external and prospective validation remain mandatory"],
        ],
        widths=[1.75, 4.95],
        font_size=8.8,
    )


def add_problem_section(doc: Document) -> None:
    add_page_break(doc)
    doc.add_heading("1 Problem Understanding and Framing", level=1)
    doc.add_heading("Clinical and Operational Problem", level=2)
    add_body(
        doc,
        "Intradialytic hypotension can interrupt treatment and require rapid reassessment. Routine haemodialysis already generates demographic, session, machine, and blood-pressure data. DIAL-ALERT asks whether those existing data can support early risk ranking so that limited clinical attention is directed toward sessions with the greatest observed risk."
    )
    add_body(
        doc,
        "The intended user is a dialysis clinician or nurse. The proposed output is a probability and a capacity-aware alert, not a treatment order. The model should supplement observation and judgment; it should never suppress routine monitoring or direct an intervention without a clinician."
    )

    doc.add_heading("Data Science Task", level=2)
    add_table(
        doc,
        ["Element", "Prespecified definition"],
        [
            ["Task", "Supervised binary classification"],
            ["Analysis unit", "One eligible patient-day haemodialysis session"],
            ["Prediction time", "Earliest valid active-dialysis observation during minutes 0 to 30"],
            ["Target", "Any later systolic blood-pressure measurement below 90 mmHg"],
            ["Predictors", "Index-time variables and patient history derived only from earlier sessions"],
            ["Exclusions", "Index SBP below 90 mmHg, inadequate later measurements, or last observed dialysis minute below 120"],
        ],
        widths=[1.65, 5.05],
        font_size=8.8,
    )

    doc.add_heading("Success Measures", level=2)
    add_table(
        doc,
        ["Type", "Measure", "Reason"],
        [
            ["Primary technical", "Average precision", "Assesses event ranking under 8.49% prevalence and is more informative than accuracy alone"],
            ["Secondary technical", "ROC AUC, Brier score, and log loss", "Assesses ranking and probability quality"],
            ["Threshold performance", "Sensitivity, specificity, precision, and F1", "Quantifies misses and false alerts at the operating threshold"],
            ["Operations and business", "Capacity metrics, review time, intervention rate, outcomes, and net cost", "Measures workload and future real-world value without assuming savings"],
            ["Equity", "Selection, error, calibration, and disparity measures", "Checks whether overall performance hides subgroup differences"],
        ],
        widths=[1.25, 1.7, 3.75],
        font_size=8.2,
        vertical_margin=65,
    )


def add_data_section(doc: Document, cohort: dict) -> None:
    add_page_break(doc)
    doc.add_heading("2 Data Collection and Understanding", level=1)
    doc.add_heading("Provenance and License", level=2)
    add_body(
        doc,
        "HEMOBP Version 3 was selected because it links patient characteristics, session summaries, and time-stamped dialysis-machine measurements. The public release is versioned on Figshare, has stable checksums, and is licensed under CC BY 4.0. The peer-reviewed data descriptor documents the collection context. [1,2]"
    )
    add_table(
        doc,
        ["Source file", "Grain", "Rows", "Verified MD5"],
        [
            ["idp.csv", "Patient", "1,072", "31d269f59bf4acc719f392ad9164d08a"],
            ["d1.csv", "Dialysis session", "165,986", "7db3b48bae2c73e3ecd731dbcd452233"],
            ["vip.csv", "Time-stamped monitor record", "4,366,298", "ff1589610cf7aa3b3e01ba979b6bdc44"],
        ],
        widths=[1.05, 1.6, 1.0, 3.05],
        font_size=8.1,
    )
    add_figure(
        doc,
        ARTIFACTS / "data_architecture.png",
        "Figure 1. HEMOBP source tables and the DIAL-ALERT session-level analytic table.",
        width=6.55,
        alt_text="Source tables linked by patient identifier and date into a session-level analytic table",
    )

    doc.add_heading("Cohort Construction", level=2)
    add_body(
        doc,
        "The source tables are joined on patient identifier and normalized session date. The earliest valid active-dialysis reading in the first 30 minutes defines the index record. Later readings define outcomes. Sessions require at least two later measurement minutes and an observation at or beyond dialysis minute 120 (not 120 minutes after prediction)."
    )
    add_table(
        doc,
        ["Cohort stage", "Count"],
        [
            ["Valid unique patient days", f"{cohort['d1_valid_unique_patient_days']:,}"],
            ["Patient days linked to monitoring", f"{cohort['linked_patient_days']:,}"],
            ["Sessions with an index record", f"{cohort['patient_days_with_index_record']:,}"],
            ["Final eligible sessions", f"{cohort['final_eligible_sessions']:,}"],
            ["Unique analytic patients", f"{cohort['unique_patients_final']:,}"],
            ["Primary events", f"{cohort['primary_idh_events']:,}"],
        ],
        widths=[4.8, 1.9],
        font_size=8.6,
        vertical_margin=55,
    )
    add_figure(
        doc,
        ARTIFACTS / "cohort_flow.png",
        "Figure 2. Deterministic cohort flow from valid session dates to the final eligible cohort.",
        width=6.25,
        alt_text="Horizontal bars showing counts at each cohort construction stage",
    )

    doc.add_heading("Data Quality and Dictionary", level=2)
    add_body(
        doc,
        "Dates, linkage, duplicate patient-days, exact monitor duplicates, documented physiologic bounds, and follow-up completeness were audited before modelling. Missing values occur mainly in dialysis vintage and first-session history features. Plausible extremes were retained; removing unusual but possible dialysis sessions could hide clinically important risk."
    )
    add_body(
        doc,
        "The complete machine-readable dictionary in docs/data_dictionary.csv and docs/data_dictionary.md records each variable's source, data type, unit, allowed values, missingness, modelling role, and prediction-time availability. A separate feature-engineering dictionary documents the derived variables."
    )
    add_bullets(
        doc,
        [
            "The publication reports 1,075 outpatients, while the released patient file contains 1,072 unique rows; the analysis uses the actual public files and records this discrepancy.",
            "Symptoms, interventions, medication timing, laboratory data, and broad comorbidity data are unavailable.",
            "The target is therefore described as BP-defined IDH rather than a complete symptomatic clinical diagnosis.",
        ],
    )


def add_eda_section(doc: Document) -> None:
    add_page_break(doc)
    doc.add_heading("3 Preprocessing EDA and Feature Engineering", level=1)
    doc.add_heading("Cleaning and Leakage Control", level=2)
    add_table(
        doc,
        ["Step", "Implementation and justification"],
        [
            ["Missing values", "Median imputation for numeric variables and mode imputation for categorical variables inside each training fold"],
            ["Duplicates", "Stable removal of duplicate patient-days and exact duplicate monitor records"],
            ["Outliers", "Documented monitoring bounds remove impossible values; plausible extremes remain for audit and sensitivity analysis"],
            ["Encoding", "One-hot encoding for recorded sex and diabetes within the fitted pipeline"],
            ["Scaling", "Standardization for linear, L1, and PCA candidates; tree models retain original numeric scale"],
            ["Temporal boundary", "Only index-time variables and shifted prior-session history enter the feature matrix"],
            ["Split boundary", "Imputation, encoding, scaling, selection, PCA, and fitting occur within grouped training folds"],
        ],
        widths=[1.45, 5.25],
        font_size=8.6,
    )

    doc.add_heading("Domain Derived Features", level=2)
    add_body(
        doc,
        "Derived variables translate raw dialysis records into quantities that can be interpreted clinically: fluid excess in kilograms and percent of dry weight, ultrafiltration intensity in mL/kg/h, mean arterial pressure, pulse pressure, dialysis vintage, and longitudinal history of prior IDH and nadir SBP. The count of prior sessions is log transformed for modelling."
    )
    add_figure(
        doc,
        ARTIFACTS / "binned_relationships.png",
        "Figure 3. Observed outcome prevalence across bins of selected predictors. These are associations, not causal effects.",
        width=6.35,
        alt_text="Binned relationships between key predictors and observed event prevalence",
    )

    doc.add_heading("Feature Selection and Dimensionality Reduction", level=2)
    add_body(
        doc,
        "Embedded L1 logistic regression retained 12 of 24 transformed predictors. This candidate tested whether sparse selection could simplify the model without using test outcomes. PCA was fitted after numeric imputation and standardization; 11 components explained 87.6% of numeric-feature variance, exceeding the prespecified 85% threshold."
    )
    add_figure(
        doc,
        ARTIFACTS / "pca_variance.png",
        "Figure 4. Cumulative explained variance of the numeric PCA candidate.",
        width=5.9,
        alt_text="Cumulative PCA explained variance reaching 87.6 percent with 11 components",
    )
    add_body(
        doc,
        "PCA was retained as a required dimensionality-reduction experiment rather than assumed to be beneficial. Its logistic-regression candidate had lower grouped-CV average precision than the selected random forest, showing that a lower-dimensional representation did not improve the primary objective."
    )


def add_model_section(doc: Document, metrics: dict) -> None:
    add_page_break(doc)
    doc.add_heading("4 Model Implementation and Comparison", level=1)
    doc.add_heading("Candidate Models and Validation", level=2)
    add_body(
        doc,
        "The comparison includes a prevalence baseline, logistic regression, embedded L1 feature selection, decision tree, PCA logistic regression, random forest, and histogram gradient boosting. Deep learning was not added because the inputs are structured tabular variables, there are only 830 independent patients, and calibration and interpretability are more important than model complexity."
    )
    add_figure(
        doc,
        ARTIFACTS / "step4_model_workflow.png",
        "Figure 5. End-to-end model workflow from eligible session predictors to probability and alert output.",
        width=6.6,
        alt_text="Model workflow showing predictors, preprocessing, random forest, probability, and alert rule",
    )
    add_table(
        doc,
        ["Partition", "Patients", "Sessions", "Events", "Prevalence"],
        [
            ["Training", "496", "64,053", "5,440", "8.49%"],
            ["Validation", "164", "21,351", "1,813", "8.49%"],
            ["Test", "170", "21,354", "1,814", "8.49%"],
        ],
        widths=[1.25, 1.15, 1.4, 1.35, 1.55],
        font_size=9.0,
    )

    doc.add_heading("Model Selection", level=2)
    add_body(
        doc,
        "Histogram gradient boosting had the highest grouped-CV average precision at 0.446, while random forest achieved 0.440. Because the two were within the prespecified 0.01 tolerance, validation Brier score determined the choice. Random forest had a validation Brier score of 0.058 compared with 0.134 for histogram gradient boosting and was selected."
    )
    add_figure(
        doc,
        ARTIFACTS / "step4_model_comparison.png",
        "Figure 6. Candidate model comparison using grouped-CV average precision and validation Brier score.",
        width=6.6,
        alt_text="Comparison of seven models on grouped cross-validation average precision and validation Brier score",
    )

    doc.add_heading("Locked Test Performance", level=2)
    ci = pd.read_csv(ARTIFACTS / "test_metric_confidence_intervals.csv").set_index("metric")
    metric_labels = [
        ("average_precision", "Average precision"),
        ("roc_auc", "ROC AUC"),
        ("brier_score", "Brier score"),
        ("sensitivity", "Sensitivity"),
        ("specificity", "Specificity"),
        ("precision", "Precision"),
    ]
    rows = []
    for key, label in metric_labels:
        row = ci.loc[key]
        rows.append([
            label,
            f"{row['estimate']:.3f}",
            f"{row['ci_lower_95']:.3f} to {row['ci_upper_95']:.3f}",
        ])
    add_table(
        doc,
        ["Metric", "Estimate", "Patient-bootstrap 95% interval"],
        rows,
        widths=[2.3, 1.4, 3.0],
        font_size=8.9,
    )
    add_figure(
        doc,
        ARTIFACTS / "roc_pr_curves.png",
        "Figure 7. Locked test ROC and precision-recall curves.",
        width=6.35,
        alt_text="ROC and precision-recall curves for the selected model on the locked test set",
    )
    add_figure(
        doc,
        ARTIFACTS / "step4_calibration_plot.png",
        "Figure 8. Test calibration curve. Sparse high-risk bins require cautious interpretation.",
        width=6.3,
        alt_text="Observed event rate plotted against predicted probability by calibration bin",
    )

    doc.add_heading("Operational Capacity", level=2)
    add_body(
        doc,
        f"At the validation-selected probability threshold of {metrics['threshold']:.3f}, the model alerts on {100 * (metrics['tp'] + metrics['fp']) / 21354:.1f}% of test sessions. A separate ranking analysis fixes the review capacity at 20% of sessions. Under that constraint, the highest-risk group contains {100 * metrics['recall_at_capacity']:.1f}% of events, but {metrics['false_alerts_per_100_sessions']:.1f} non-event sessions per 100 total sessions would also be reviewed."
    )
    add_figure(
        doc,
        ARTIFACTS / "capacity_curve.png",
        "Figure 9. Recall, precision, lift, and false-alert workload across review capacity levels.",
        width=6.05,
        alt_text="Capacity curve showing tradeoffs between event capture and false-alert burden",
    )


def add_ethics_section(doc: Document) -> None:
    doc.add_heading("5 Critical Thinking Ethical AI and Bias Auditing", level=1)
    doc.add_heading("Model Explanations", level=2)
    add_body(
        doc,
        "Permutation importance, aggregate approximate interventional SHAP values, synthetic SHAP and LIME local explanations, and dependence curves across synthetic reference profiles provide complementary views of the fitted model without publishing patient-level explanation records. Prior-session nadir SBP and prior IDH rate were the strongest permutation features. Explanation outputs describe learned associations; they do not identify safe treatment targets or causal mechanisms. [5-7]"
    )
    add_figure(
        doc,
        ARTIFACTS / "step5_shap_summary.png",
        "Figure 10. Aggregate approximate interventional SHAP importance for the locked model.",
        width=6.1,
        alt_text="Aggregate SHAP bar chart showing mean absolute feature contributions",
    )
    add_figure(
        doc,
        ARTIFACTS / "step5_pdp_ice.png",
        "Figure 11. Dependence and conditional-effect patterns across synthetic reference profiles.",
        width=6.55,
        alt_text="Dependence and conditional-effect curves for key predictors using synthetic reference profiles",
    )

    doc.add_heading("Limitations and Robustness", level=2)
    add_table(
        doc,
        ["Risk", "Evidence", "Response"],
        [
            ["Class imbalance", "Test prevalence is 8.49%", "Use average precision, precision, sensitivity, calibration, and workload rather than accuracy alone"],
            ["Patient leakage", "No patient overlaps training, validation, and test", "Retain grouped splitting in every future evaluation"],
            ["Temporal leakage", "Outcomes occur after the index time; history is shifted", "Preserve the feature timestamp contract in production"],
            ["Optimism", "Grouped-CV AP 0.440, validation AP 0.470, test AP 0.395", "Report patient-bootstrap uncertainty and require external validation"],
            ["Repeated sessions", "21,354 test sessions come from 170 patients", "Use patient-cluster bootstrap and patient-equal-weight sensitivity analyses"],
            ["Outcome validity", "Symptoms and interventions are unavailable", "Describe the endpoint as BP-defined IDH"],
        ],
        widths=[1.25, 2.35, 3.1],
        font_size=8.2,
    )

    doc.add_heading("Fairness Audit", level=2)
    disparities = pd.read_csv(ARTIFACTS / "step5_disparities.csv")
    disparity_rows = []
    for _, row in disparities.iterrows():
        disparity_rows.append([
            row["attribute"],
            f"{row['disparate_impact_ratio']:.3f}",
            f"{row['equal_opportunity_difference']:.3f}",
            f"{row['false_positive_rate_difference']:.3f}",
            f"{row['brier_score_difference']:.3f}",
        ])
    add_table(
        doc,
        ["Attribute", "Disparate impact ratio", "Sensitivity gap", "FPR gap", "Brier gap"],
        disparity_rows,
        widths=[1.45, 1.55, 1.25, 1.15, 1.3],
        font_size=8.3,
    )
    add_body(
        doc,
        "Recorded sex and age are the only available protected-characteristic audits. Diabetes is a clinical subgroup and is not used as a socioeconomic proxy. Race, ethnicity, socioeconomic status, and gender identity are absent, so fairness across those groups cannot be assessed. The 0.80 disparate-impact ratio is treated as a descriptive screening heuristic, not a fairness certificate. [8]"
    )
    add_figure(
        doc,
        ARTIFACTS / "step5_fairness_overview.png",
        "Figure 12. Selection and error metrics across available groups.",
        width=6.55,
        alt_text="Fairness overview showing selection, sensitivity, false-positive rate, and precision by group",
    )

    doc.add_heading("Mitigation Experiments", level=2)
    mitigation = pd.read_csv(ARTIFACTS / "step5_mitigation_comparison.csv")
    mitigation_rows = []
    for _, row in mitigation.iterrows():
        mitigation_rows.append([
            row["strategy"],
            f"{row['average_precision']:.3f}",
            f"{row['precision']:.3f}",
            f"{row['sex_disparate_impact_ratio']:.3f}",
            f"{row['sex_equalized_odds_difference']:.3f}",
        ])
    add_table(
        doc,
        ["Strategy", "Average precision", "Precision", "Sex impact ratio", "Equalized odds gap"],
        mitigation_rows,
        widths=[2.0, 1.3, 1.1, 1.25, 1.25],
        font_size=8.2,
    )
    add_body(
        doc,
        "Outcome-by-sex training weights improved the recorded-sex impact ratio from 0.705 to 0.764 and reduced the equalized-odds gap from 0.076 to 0.058, while average precision changed from 0.395 to 0.394. Sex-specific thresholds worsened held-out fairness and precision and are rejected. Reweighting remains a candidate for external testing, not an automatic deployment choice."
    )
    add_figure(
        doc,
        ARTIFACTS / "step5_mitigation_tradeoff.png",
        "Figure 13. Performance and fairness tradeoffs across mitigation strategies.",
        width=6.45,
        alt_text="Comparison of original, reweighted, and sex-specific threshold strategies",
    )


def add_business_section(doc: Document, metrics: dict) -> None:
    add_page_break(doc)
    doc.add_heading("6 Final Communication and Business Translation", level=1)
    doc.add_heading("Proposed Workflow", level=2)
    add_body(
        doc,
        "A future silent-mode pilot would calculate risk after the index observation, place the highest-risk sessions on a review list, and compare predictions with subsequent BP measurements without changing care. Only after prespecified safety and workload review should a later phase test whether presenting alerts changes clinician action or patient outcomes."
    )
    add_table(
        doc,
        ["Per 100 eligible sessions at 20% capacity", "Retrospective estimate"],
        [
            ["Sessions selected for review", "20.0"],
            ["Observed events in the cohort", "8.5"],
            ["Observed events concentrated in the selected group", "About 5.9"],
            ["Observed events outside the selected group", "About 2.6"],
            ["Selected sessions without the event", f"About {metrics['false_alerts_per_100_sessions']:.1f}"],
        ],
        widths=[4.55, 2.15],
        font_size=8.8,
    )
    add_body(
        doc,
        "These values describe risk concentration in retrospective data. They are not prevented events. Every selected session still requires clinical interpretation, and some unselected sessions will experience the outcome."
    )

    doc.add_heading("ROI Measurement Framework", level=2)
    add_body(
        doc,
        "The pilot should measure net value rather than assume it. A transparent calculation is: net pilot value equals the number of events causally prevented multiplied by the local value of avoiding one event, less review time, implementation, maintenance, and monitoring costs. The prevention rate and local cost inputs are currently unknown."
    )
    add_table(
        doc,
        ["Input", "Pilot measurement"],
        [
            ["Alert workload", "Alerts per 100 sessions and median review minutes"],
            ["Clinical response", "Proportion reviewed, action taken, action type, and override reason"],
            ["Patient outcome", "BP-defined events, symptoms, interventions, interrupted treatment, and recovery time"],
            ["Safety", "Missed events, delayed treatment, inappropriate action, and incident reports"],
            ["Equity", "Selection, sensitivity, false-positive rate, precision, and calibration by governed groups"],
            ["Cost", "Staff time, implementation, infrastructure, training, maintenance, and event-related utilization"],
        ],
        widths=[1.6, 5.1],
        font_size=8.6,
    )

    doc.add_heading("Governance and Pilot Gates", level=2)
    add_table(
        doc,
        ["Gate", "Minimum evidence before progression"],
        [
            ["External validity", "Acceptable discrimination and calibration in a different centre or later time period"],
            ["Operational fit", "Review volume and response time remain within agreed staffing capacity"],
            ["Safety", "No unacceptable harm signal, with defined clinician override and escalation"],
            ["Fairness", "Prespecified subgroup metrics and uncertainty reviewed by governance"],
            ["Human factors", "Clinicians understand the score, avoid automation bias, and document decisions"],
            ["Monitoring", "Data quality, drift, calibration, alert burden, incidents, and rollback are operational"],
        ],
        widths=[1.6, 5.1],
        font_size=8.6,
    )
    add_body(
        doc,
        "The final communication package includes a sixteen-slide technical deck for peers and an eleven-slide business deck for executives. Both decks state the evidence boundary and recommend validation rather than deployment."
    )


def add_repository_section(doc: Document) -> None:
    add_page_break(doc)
    doc.add_heading("7 Reproducibility and GitHub Repository", level=1)
    add_body(
        doc,
        "The public repository is organized as an open-source project. Raw and processed data are excluded from Git because the source is large and patient-level. A version-pinned acquisition script downloads HEMOBP Version 3 and verifies MD5 checksums. The selected compressed predictor is retained so that inference can be tested without retraining."
    )
    add_table(
        doc,
        ["Repository path", "Contents"],
        [
            ["src", "Acquisition, cohort construction, EDA, training, audit, inference, and report builders"],
            ["notebooks", "Guided end-to-end workflow"],
            ["data", "Acquisition instructions and empty tracked directories"],
            ["models", "Selected predictor, threshold specification, and manifest"],
            ["artifacts", "Metrics, uncertainty intervals, audit tables, and figures"],
            ["docs", "Data dictionary, preprocessing specification, and model card"],
            ["reports", "Final report, component reports, and technical and business presentations"],
            ["tests", "Contract and saved-model inference tests"],
            [".github workflows", "Automated tests on each push and pull request"],
        ],
        widths=[1.75, 4.95],
        font_size=8.6,
    )
    commands = [
        "python -m pip install -r requirements.txt",
        "python src/download_data.py",
        "python src/build_session_dataset.py --raw-dir data/raw --output-dir data/processed",
        "python src/generate_step2_assets.py",
        "python src/train_evaluate.py --data data/processed/hemobp_session_level.csv.gz --config configs/model_config.json --artifacts artifacts --models models",
        "python src/generate_eda.py",
        "python src/generate_step4_assets.py",
        "python src/audit_bias_fairness.py",
        "python src/create_final_report.py",
        "python -m pytest -q",
    ]
    p = doc.add_paragraph()
    p.paragraph_format.right_indent = Inches(0.1)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_together = True
    heading_run = p.add_run("Reproduction Commands")
    heading_run.bold = True
    set_font(heading_run, size=12)
    heading_run.add_break()
    run = p.add_run("\n".join(commands))
    set_font(run, name="Courier New", size=7.4)
    add_body(
        doc,
        "The model manifest records the data and configuration SHA-256 hashes, Python and package versions, selected model, random seeds, and hashes of saved artifacts. Continuous integration runs repository and prediction smoke tests. The notebook presents the same sequence with explanatory text."
    )


def add_conclusion(doc: Document) -> None:
    add_page_break(doc)
    doc.add_heading("8 Conclusions and Next Steps", level=1)
    add_body(
        doc,
        "DIAL-ALERT demonstrates the complete machine-learning lifecycle on a real, industry-relevant dataset. The project frames an auditable clinical question, constructs a leakage-controlled cohort, compares multiple models, evaluates an untouched patient-disjoint test set, explains model behavior, audits available groups, tests mitigation, and translates performance into operational workload."
    )
    add_body(
        doc,
        "The main finding is useful but limited: the random forest concentrates 69.1% of observed BP-defined events within the highest-risk 20% of sessions. This supports further validation. It does not demonstrate prevented events, safer dialysis, reduced cost, or transportability to another centre."
    )
    doc.add_heading("Recommended Sequence", level=2)
    add_bullets(
        doc,
        [
            "Validate the locked model on an external centre and a later time period without retuning.",
            "Collect symptoms, interventions, medication timing, broader comorbidities, and governed demographic attributes.",
            "Run a prospective silent-mode study to measure data availability, calibration, workload, and drift.",
            "Prespecify subgroup and safety thresholds, clinician override, incident response, and rollback rules.",
            "Only then test whether visible alerts change actions and improve patient outcomes in a controlled implementation study.",
        ],
    )
    doc.add_heading("Final Decision", level=2)
    add_body(
        doc,
        "Continue DIAL-ALERT as a validation-stage clinical decision-support research project. Do not deploy it for patient-care decisions at this stage."
    )


def add_references(doc: Document) -> None:
    add_page_break(doc)
    doc.add_heading("References", level=1)
    references = [
        "1. Lin CJ, Chen YY, Pan CF, Wu VC, Wu CJ. Dataset supporting blood pressure prediction for the management of chronic hemodialysis. Scientific Data. 2019;6:313. https://doi.org/10.1038/s41597-019-0319-8",
        "2. Chien CY. HEMOBP Version 3. Figshare dataset. 2019. CC BY 4.0. https://doi.org/10.6084/m9.figshare.6260654.v3",
        "3. Breiman L. Random forests. Machine Learning. 2001;45:5-32. https://doi.org/10.1023/A:1010933404324",
        "4. Saito T, Rehmsmeier M. The precision-recall plot is more informative than the ROC plot when evaluating binary classifiers on imbalanced datasets. PLoS ONE. 2015;10:e0118432. https://doi.org/10.1371/journal.pone.0118432",
        "5. Lundberg SM, Lee SI. A unified approach to interpreting model predictions. Advances in Neural Information Processing Systems. 2017;30. https://papers.nips.cc/paper/7062-a-unified-approach-to-interpreting-model-predictions",
        "6. Ribeiro MT, Singh S, Guestrin C. \"Why Should I Trust You?\": Explaining the Predictions of Any Classifier. Proceedings of KDD. 2016:1135-1144. https://doi.org/10.1145/2939672.2939778",
        "7. Goldstein A, Kapelner A, Bleich J, Pitkin E. Peeking Inside the Black Box: Visualizing Statistical Learning With Plots of Individual Conditional Expectation. Journal of Computational and Graphical Statistics. 2015;24:44-65. https://doi.org/10.1080/10618600.2014.907095",
        "8. Hardt M, Price E, Srebro N. Equality of opportunity in supervised learning. Advances in Neural Information Processing Systems. 2016;29. https://arxiv.org/abs/1610.02413",
        "9. World Health Organization. Ethics and governance of artificial intelligence for health. 2021. https://www.who.int/publications/i/item/9789240029200",
        "10. Collins GS, Moons KGM, Dhiman P, Riley RD, Beam AL, Van Calster B, et al. TRIPOD+AI statement: updated guidance for reporting clinical prediction models that use regression or machine learning methods. BMJ. 2024;385:e078378. https://doi.org/10.1136/bmj-2023-078378",
        "11. Flythe JE, Xue H, Lynch KE, Curhan GC, Brunelli SM. Association of mortality risk with various definitions of intradialytic hypotension. Journal of the American Society of Nephrology. 2015;26:724-734. https://doi.org/10.1681/ASN.2014020222",
    ]
    for reference in references:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(-0.2)
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(reference)
        set_font(run, size=8.7)


def add_rubric_appendix(doc: Document) -> None:
    add_page_break(doc)
    doc.add_heading("Appendix Project Evidence Map", level=1)
    add_table(
        doc,
        ["Project step", "Evidence"],
        [
            ["1 Problem Understanding and Framing", "Clear clinical context, supervised task, target, technical metrics, operational KPIs, and evidence boundary"],
            ["2 Data Collection and Understanding", "Peer-reviewed public dataset, license, checksums, cohort audit, complete dictionaries, and limitations"],
            ["3 Preprocessing EDA and Feature Engineering", "Nulls, duplicates, outliers, encoding, scaling, domain features, visual EDA, L1 selection, PCA, and explainability"],
            ["4 Model Implementation and Comparison", "Seven configurations, grouped tuning, locked selection rule, saved models, test metrics, intervals, and workload analysis"],
            ["5 Critical Thinking Ethical AI and Bias Auditing", "SHAP, local surrogate, PDP and ICE, leakage and overfitting analysis, fairness metrics, uncertainty, and mitigation tests"],
            ["6 Final Presentation and Communication", "Separate technical and business decks with notes, citations, ROI framework, risks, strategy, and pilot recommendation"],
            ["7 GitHub Profile and Upload", "Open-source structure, README, license, citation file, data acquisition, notebooks, tests, CI, final report, and reproducible commands"],
        ],
        widths=[2.35, 4.35],
        font_size=8.3,
    )


def make_report() -> Path:
    cohort = load_json(ROOT / "data/processed/cohort_flow.json")
    metrics = load_json(ARTIFACTS / "final_test_metrics.json")["test_metrics_calibrated"]

    doc = Document()
    configure_document(doc)
    replace_footer(doc)
    doc.core_properties.title = "DIAL ALERT Capstone Final Report"
    doc.core_properties.subject = "End to end machine learning capstone"
    doc.core_properties.author = "Franklin B. Guillano"
    doc.core_properties.keywords = "haemodialysis, intradialytic hypotension, machine learning, fairness"
    doc.core_properties.comments = "DIAL-ALERT academic capstone"

    add_title_page(doc, cohort, metrics)
    add_report_map(doc)
    add_executive_summary(doc, cohort, metrics)
    add_problem_section(doc)
    add_data_section(doc, cohort)
    add_eda_section(doc)
    add_model_section(doc, metrics)
    add_ethics_section(doc)
    add_business_section(doc, metrics)
    add_repository_section(doc)
    add_conclusion(doc)
    add_references(doc)
    add_rubric_appendix(doc)

    from revise_submission_documents import append_notes
    append_notes(doc)
    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT}")
    return OUTPUT


if __name__ == "__main__":
    make_report()
